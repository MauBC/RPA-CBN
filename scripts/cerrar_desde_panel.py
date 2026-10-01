from pathlib import Path
import sys
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

ROOT_DIR = Path(__file__).resolve().parents[1]

if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from src.browser.session import URL_SISTEMA, abrir_contexto
from src.browser.auth import esperar_login_manual
from src.excel.reader import leer_ordenes_excel
from src.excel.result_writer import guardar_resultados_ordenes
from src.excel.state_manager import obtener_fila_excel_por_id, actualizar_estado_orden
from src.flows.cerrar_cotizacion import cerrar_cotizacion_desde_panel
from src.utils.debug import preparar_carpetas, guardar_evidencia_error
from src.utils.user_errors import construir_mensaje_usuario, construir_detalle_tecnico


RUTA_EXCEL = ROOT_DIR / "data" / "DATA.xlsx"
HOJA_ORDENES = "Ordenes"
COL_ID_ORDEN = "ID_ORDEN"
COL_ESTADO_RPA = "ESTADO_RPA"


def _normalizar_id(valor) -> str:
    if valor is None:
        return ""

    try:
        if pd.isna(valor):
            return ""
    except Exception:
        pass

    texto = str(valor).strip()

    if texto.endswith(".0"):
        texto = texto[:-2]

    return texto


def _estado_int(valor) -> int | None:
    if valor is None:
        return None

    try:
        if pd.isna(valor):
            return None
    except Exception:
        pass

    texto = str(valor).strip()

    if texto.endswith(".0"):
        texto = texto[:-2]

    try:
        return int(texto)
    except Exception:
        return None


def _headers_ws(ws) -> dict[str, int]:
    headers = {}

    for cell in ws[1]:
        nombre = str(cell.value or "").strip().upper()

        if nombre:
            headers[nombre] = cell.column

    return headers


def validar_ids_unicos_y_estado(ruta_excel: Path) -> None:
    wb = load_workbook(ruta_excel)

    if HOJA_ORDENES not in wb.sheetnames:
        raise RuntimeError(f"No existe la hoja {HOJA_ORDENES}.")

    ws = wb[HOJA_ORDENES]
    headers = _headers_ws(ws)

    if COL_ID_ORDEN not in headers:
        raise RuntimeError(f"No existe la columna {COL_ID_ORDEN}.")

    if COL_ESTADO_RPA not in headers:
        raise RuntimeError(f"No existe la columna {COL_ESTADO_RPA}.")

    col_id = headers[COL_ID_ORDEN]

    vistos = {}
    duplicados = []

    for fila in range(2, ws.max_row + 1):
        id_orden = _normalizar_id(ws.cell(row=fila, column=col_id).value)

        if not id_orden:
            continue

        if id_orden in vistos:
            duplicados.append(f"ID_ORDEN={id_orden} filas Excel {vistos[id_orden]} y {fila}")
        else:
            vistos[id_orden] = fila

    if duplicados:
        raise RuntimeError(
            "Hay ID_ORDEN repetidos. Corrige antes de ejecutar:\n"
            + "\n".join(duplicados)
        )


def obtener_ids_para_cerrar(ruta_excel: Path) -> list[str]:
    """
    Este checkpoint procesa órdenes pendientes con ESTADO_RPA = 0.
    Si cierra bien, cambia a 1.
    Si falla, cambia a 2.
    """
    validar_ids_unicos_y_estado(ruta_excel)

    wb = load_workbook(ruta_excel, data_only=True)
    ws = wb[HOJA_ORDENES]
    headers = _headers_ws(ws)

    col_id = headers[COL_ID_ORDEN]
    col_estado = headers[COL_ESTADO_RPA]

    ids = []

    for fila in range(2, ws.max_row + 1):
        id_orden = _normalizar_id(ws.cell(row=fila, column=col_id).value)
        estado = _estado_int(ws.cell(row=fila, column=col_estado).value)

        if id_orden and estado == 0:
            ids.append(id_orden)

    return ids


def crear_excel_trabajo_cierre(ruta_excel: Path) -> tuple[Path, list[str]]:
    ids = obtener_ids_para_cerrar(ruta_excel)
    ids_set = set(ids)

    ruta_trabajo = ruta_excel.parent / f"_{ruta_excel.stem}_cierre_ESTADO_RPA{ruta_excel.suffix}"

    hojas = pd.read_excel(ruta_excel, sheet_name=None, dtype=object)
    hojas_filtradas = {}

    for nombre_hoja, df in hojas.items():
        df_filtrado = df.copy()

        if COL_ID_ORDEN in df_filtrado.columns:
            serie_ids = df_filtrado[COL_ID_ORDEN].map(_normalizar_id)
            df_filtrado = df_filtrado[serie_ids.isin(ids_set)].copy()

        hojas_filtradas[nombre_hoja] = df_filtrado

    with pd.ExcelWriter(ruta_trabajo, engine="openpyxl") as writer:
        for nombre_hoja, df in hojas_filtradas.items():
            df.to_excel(writer, sheet_name=nombre_hoja[:31], index=False)

    return ruta_trabajo, ids


def construir_resultado_ok(indice: int, orden, info_cierre) -> dict:
    return {
        "NRO": indice,
        "ID_ORDEN": orden.id_orden,
        "FILA_EXCEL": obtener_fila_excel_por_id(RUTA_EXCEL, orden.id_orden) or "",
        "ESTADO": "OK_CIERRE_COTIZACION",
        "N_COTIZACION": getattr(info_cierre, "numero_cotizacion", ""),
        "ESTADO_CBN": getattr(info_cierre, "estado_cbn", ""),
        "MENSAJE_USUARIO": getattr(info_cierre, "mensaje", ""),
        "PASO_ERROR": "",
        "DETALLE_TECNICO": "",
        "LOG": "",
        "SCREENSHOT": "",
        "HTML": "",
        "FECHA_HORA": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


def construir_resultado_error(indice: int, orden, error: Exception, rutas: dict) -> dict:
    return {
        "NRO": indice,
        "ID_ORDEN": getattr(orden, "id_orden", ""),
        "FILA_EXCEL": obtener_fila_excel_por_id(RUTA_EXCEL, getattr(orden, "id_orden", "")) or "",
        "ESTADO": "ERROR_CIERRE_COTIZACION",
        "N_COTIZACION": "",
        "ESTADO_CBN": "",
        "MENSAJE_USUARIO": construir_mensaje_usuario(error, orden),
        "PASO_ERROR": getattr(error, "paso", "cerrar_cotizacion"),
        "DETALLE_TECNICO": construir_detalle_tecnico(error),
        "LOG": str(rutas.get("log", "")),
        "SCREENSHOT": str(rutas.get("screenshot", "")),
        "HTML": str(rutas.get("html", "")),
        "FECHA_HORA": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


def main() -> None:
    preparar_carpetas()

    playwright = None
    contexto = None
    page = None
    resultados = []

    try:
        ruta_excel_trabajo, ids = crear_excel_trabajo_cierre(RUTA_EXCEL)

        if not ids:
            print("No hay órdenes para cerrar.")
            print("Este checkpoint procesa órdenes con ESTADO_RPA = 0.")
            return

        ordenes = leer_ordenes_excel(ruta_excel_trabajo)

        if not ordenes:
            raise RuntimeError("No hay órdenes válidas para cerrar.")

        print("============================================================")
        print("CHECKPOINT: Cerrar cotización desde Panel")
        print("Este script NO manda a cotizar.")
        print("Procesa órdenes con ESTADO_RPA = 0.")
        print("Si cierra bien, cambia ESTADO_RPA a 1.")
        print("Si falla, cambia ESTADO_RPA a 2.")
        print("============================================================")
        print(f"Órdenes a cerrar: {len(ordenes)}")

        playwright, contexto, page = abrir_contexto()

        page.goto(
            URL_SISTEMA,
            wait_until="domcontentloaded",
            timeout=60_000,
        )

        esperar_login_manual(page)

        total = len(ordenes)

        for indice, orden in enumerate(ordenes, start=1):
            print("")
            print("------------------------------------------------------------")
            print(f"Cerrando orden {indice}/{total}")
            print(f"ID_ORDEN: {orden.id_orden}")
            print(f"Fila Excel: {obtener_fila_excel_por_id(RUTA_EXCEL, orden.id_orden)}")
            print(f"Nombre de envío: {orden.texto}")
            print("------------------------------------------------------------")

            try:
                info_cierre = cerrar_cotizacion_desde_panel(page, orden)

                print("Checkpoint cierre OK")
                print(f"N° de Cotización: {info_cierre.numero_cotizacion}")
                print(f"URL DURACIÓN: {info_cierre.url_duration}")

                resultados.append(construir_resultado_ok(indice, orden, info_cierre))

                actualizar_estado_orden(RUTA_EXCEL, orden.id_orden, 1)
                print(f"ESTADO_RPA actualizado a 1 para ID_ORDEN={orden.id_orden}")

            except Exception as error:
                print("")
                print("Error en checkpoint cerrar cotización.")
                print(f"ID_ORDEN: {orden.id_orden}")
                print(f"Mensaje usuario: {construir_mensaje_usuario(error, orden)}")
                print(f"Detalle técnico: {error}")

                prefijo = f"cerrar_cotizacion_orden_{orden.id_orden}"
                rutas = guardar_evidencia_error(page, error, prefijo)

                if "log" in rutas:
                    print(f"Log guardado en: {rutas['log']}")
                if "screenshot" in rutas:
                    print(f"Screenshot guardado en: {rutas['screenshot']}")
                if "html" in rutas:
                    print(f"HTML guardado en: {rutas['html']}")

                resultados.append(construir_resultado_error(indice, orden, error, rutas))

                actualizar_estado_orden(RUTA_EXCEL, orden.id_orden, 2)
                print(f"ESTADO_RPA actualizado a 2 para ID_ORDEN={orden.id_orden}")

                print("Se continúa con la siguiente orden.")

        ruta_resultado = guardar_resultados_ordenes(RUTA_EXCEL, resultados)

        ok = sum(1 for item in resultados if item["ESTADO"] == "OK_CIERRE_COTIZACION")
        errores = len(resultados) - ok

        print("")
        print("============================================================")
        print("Checkpoint cerrar cotización terminado.")
        print(f"Órdenes cerradas OK: {ok}")
        print(f"Órdenes con error: {errores}")
        print(f"Resultado guardado en: {ruta_resultado}")
        print("============================================================")

    finally:
        if contexto is not None:
            try:
                contexto.close()
            except Exception:
                pass

        if playwright is not None:
            try:
                playwright.stop()
            except Exception:
                pass


if __name__ == "__main__":
    main()
