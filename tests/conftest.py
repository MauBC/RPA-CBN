from pathlib import Path

from openpyxl import Workbook


def crear_excel_estados(
    ruta: Path,
    filas: list[tuple],
    incluir_estado: bool = True,
) -> Path:
    """
    Crea un Excel mínimo para probar state_manager.

    filas:
        con estado:
            ("1001", 0)
            ("1001", 0)
            ("1002", 1)

        sin estado:
            ("1001",)
            ("1002",)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Ordenes"

    if incluir_estado:
        ws.append(["ID_ORDEN", "ESTADO_RPA"])
    else:
        ws.append(["ID_ORDEN"])

    for fila in filas:
        ws.append(list(fila))

    wb.save(ruta)
    wb.close()

    return ruta


def crear_template(
    ruta: Path,
    filas: list[tuple],
    header_a: str = "code",
    header_b: str = "quantityOrPercent",
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Template"

    ws["A1"] = header_a
    ws["B1"] = header_b

    # El RPA comienza a leer desde la fila 3.
    fila_excel = 3

    for codigo, valor in filas:
        ws.cell(row=fila_excel, column=1).value = codigo
        ws.cell(row=fila_excel, column=2).value = valor
        fila_excel += 1

    wb.save(ruta)
    wb.close()

    return ruta
