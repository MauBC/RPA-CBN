from __future__ import annotations

from openpyxl import Workbook, load_workbook
import pytest

from src.excel.concurrency_guard import (
    ExcelConflictoEdicionError,
    capturar_versiones_ordenes,
)
from src.excel.state_manager import (
    actualizar_resultado_orden,
)


def _crear_excel(
    ruta,
):
    wb = Workbook()

    try:
        ws = wb.active
        ws.title = "Ordenes"

        ws.append(
            [
                "ID_ORDEN",
                "ESTADO_RPA",
                "PROVEEDOR",
                "Valor",
                "RESUMEN",
            ]
        )

        ws.append(
            [
                "1001",
                0,
                "PROV-A",
                100,
                "",
            ]
        )

        ws.append(
            [
                "1001",
                0,
                "PROV-A",
                200,
                "",
            ]
        )

        ws.append(
            [
                "1002",
                0,
                "PROV-B",
                300,
                "",
            ]
        )

        adj = wb.create_sheet(
            "Adjuntos"
        )

        adj.append(
            [
                "ID_ORDEN",
                "Archivo",
            ]
        )

        adj.append(
            [
                "1001",
                "a.pdf",
            ]
        )

        wb.save(
            ruta
        )

    finally:
        wb.close()

    return ruta


def _version(
    ruta,
    id_orden,
):
    return capturar_versiones_ordenes(
        ruta,
        [
            id_orden,
        ],
    )[
        id_orden
    ]


def test_version_orden_sin_cambios_permite_persistir(
    tmp_path,
):
    ruta = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    version = _version(
        ruta,
        "1001",
    )

    filas = actualizar_resultado_orden(
        ruta,
        "1001",
        estado_rpa=1,
        resumen="OK",
        version_esperada=version,
    )

    assert filas == 2


def test_edicion_de_otra_orden_no_bloquea(
    tmp_path,
):
    ruta = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    version = _version(
        ruta,
        "1001",
    )

    wb = load_workbook(
        ruta
    )

    try:
        wb["Ordenes"]["C4"] = (
            "PROV-B-MODIFICADO"
        )

        wb.save(
            ruta
        )

    finally:
        wb.close()

    actualizar_resultado_orden(
        ruta,
        "1001",
        estado_rpa=1,
        resumen="OK",
        version_esperada=version,
    )

    wb = load_workbook(
        ruta,
        data_only=True,
    )

    try:
        assert (
            wb["Ordenes"]["C4"].value
            == "PROV-B-MODIFICADO"
        )

    finally:
        wb.close()


def test_edicion_de_misma_orden_bloquea_y_no_sobrescribe(
    tmp_path,
):
    ruta = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    version = _version(
        ruta,
        "1001",
    )

    wb = load_workbook(
        ruta
    )

    try:
        wb["Ordenes"]["C2"] = (
            "PROV-A-MODIFICADO"
        )

        wb.save(
            ruta
        )

    finally:
        wb.close()

    with pytest.raises(
        ExcelConflictoEdicionError,
        match="modificada",
    ):
        actualizar_resultado_orden(
            ruta,
            "1001",
            estado_rpa=1,
            resumen="OK",
            version_esperada=version,
        )

    wb = load_workbook(
        ruta,
        data_only=True,
    )

    try:
        assert (
            wb["Ordenes"]["B2"].value
            == 0
        )

        assert (
            wb["Ordenes"]["C2"].value
            == "PROV-A-MODIFICADO"
        )

    finally:
        wb.close()


def test_edicion_de_adjunto_de_misma_orden_bloquea(
    tmp_path,
):
    ruta = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    version = _version(
        ruta,
        "1001",
    )

    wb = load_workbook(
        ruta
    )

    try:
        wb["Adjuntos"]["B2"] = (
            "nuevo.pdf"
        )

        wb.save(
            ruta
        )

    finally:
        wb.close()

    with pytest.raises(
        ExcelConflictoEdicionError,
    ):
        actualizar_resultado_orden(
            ruta,
            "1001",
            estado_rpa=1,
            resumen="OK",
            version_esperada=version,
        )


def test_cambio_manual_estado_incompatible_bloquea(
    tmp_path,
):
    ruta = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    version = _version(
        ruta,
        "1001",
    )

    wb = load_workbook(
        ruta
    )

    try:
        wb["Ordenes"]["B2"] = 2
        wb["Ordenes"]["B3"] = 2

        wb.save(
            ruta
        )

    finally:
        wb.close()

    with pytest.raises(
        ExcelConflictoEdicionError,
        match="ESTADO_RPA",
    ):
        actualizar_resultado_orden(
            ruta,
            "1001",
            estado_rpa=1,
            resumen="OK",
            version_esperada=version,
        )


def test_resultado_ya_aplicado_es_idempotente(
    tmp_path,
):
    ruta = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    version = _version(
        ruta,
        "1001",
    )

    actualizar_resultado_orden(
        ruta,
        "1001",
        estado_rpa=1,
        resumen="OK",
        version_esperada=version,
    )

    # Un replay con la misma version y mismo resultado
    # debe ser seguro.
    filas = actualizar_resultado_orden(
        ruta,
        "1001",
        estado_rpa=1,
        resumen="OK",
        version_esperada=version,
    )

    assert filas == 2
