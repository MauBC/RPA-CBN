from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.excel.reader import (
    _leer_filas_template,
    _template_es_porcentaje,
)
from src.excel.snapshot_manager import (
    _nombre_snapshot_template,
    crear_snapshot_template_estable,
    sha256_archivo,
)


def _crear_template(
    ruta: Path,
    valores,
    *,
    formato_texto: bool = False,
):
    wb = Workbook()

    try:
        ws = wb.active
        ws.title = "Imputacion"

        ws["A1"] = "code"
        ws["B1"] = "quantityOrPercent"

        ws["A2"] = "Código"
        ws["B2"] = "Cantidad o porcentaje"

        codigos = [
            "51WB2KB092",
            "51AD000H07",
            "51AD000I07",
            "51AD000K07",
            "51AD000Q07",
            "51B0000Y04",
        ]

        for fila, (
            codigo,
            valor,
        ) in enumerate(
            zip(
                codigos,
                valores,
            ),
            start=3,
        ):
            ws.cell(
                row=fila,
                column=1,
            ).value = codigo

            celda = ws.cell(
                row=fila,
                column=2,
            )

            if formato_texto:
                celda.number_format = "@"

            celda.value = valor

        wb.save(
            ruta
        )

    finally:
        wb.close()

    return ruta


def _leer_celdas_b(
    ruta: Path,
):
    wb = load_workbook(
        ruta,
        data_only=False,
    )

    try:
        ws = wb[
            wb.sheetnames[0]
        ]

        return [
            (
                ws.cell(
                    fila,
                    2,
                ).value,
                ws.cell(
                    fila,
                    2,
                ).data_type,
                ws.cell(
                    fila,
                    2,
                ).number_format,
            )
            for fila in range(
                3,
                ws.max_row + 1,
            )
        ]

    finally:
        wb.close()


def test_snapshot_convierte_strings_a_numeros_reales(
    tmp_path,
):
    origen = _crear_template(
        tmp_path
        / "origen.xlsx",
        [
            "300",
            "200",
            "100",
            "500",
            "400",
            "235",
        ],
    )

    destino = (
        tmp_path
        / "snapshot.xlsx"
    )

    metadata = crear_snapshot_template_estable(
        origen,
        destino,
    )

    # El original permanece intacto.
    original = _leer_celdas_b(
        origen
    )

    assert all(
        data_type == "s"
        for _, data_type, _
        in original
    )

    snapshot = _leer_celdas_b(
        destino
    )

    assert [
        valor
        for valor, _, _
        in snapshot
    ] == [
        300,
        200,
        100,
        500,
        400,
        235,
    ]

    assert all(
        data_type == "n"
        for _, data_type, _
        in snapshot
    )

    assert (
        metadata[
            "normalizacion_quantity_or_percent"
        ][
            "normalizadas"
        ]
        == 6
    )

    assert (
        metadata[
            "sha256"
        ]
        == sha256_archivo(
            destino
        )
    )


def test_snapshot_elimina_formato_texto(
    tmp_path,
):
    origen = _crear_template(
        tmp_path
        / "texto.xlsx",
        [
            "300",
            "200",
            "100",
            "500",
            "400",
            "235",
        ],
        formato_texto=True,
    )

    destino = (
        tmp_path
        / "snapshot.xlsx"
    )

    crear_snapshot_template_estable(
        origen,
        destino,
    )

    snapshot = _leer_celdas_b(
        destino
    )

    assert all(
        data_type == "n"
        for _, data_type, _
        in snapshot
    )

    assert all(
        formato != "@"
        for _, _, formato
        in snapshot
    )


def test_snapshot_normaliza_espacios_y_decimales(
    tmp_path,
):
    origen = _crear_template(
        tmp_path
        / "decimales.xlsx",
        [
            " 300.50 ",
            "200.25",
            "100",
            "500",
            "400",
            "234.25",
        ],
    )

    destino = (
        tmp_path
        / "snapshot.xlsx"
    )

    crear_snapshot_template_estable(
        origen,
        destino,
    )

    filas = _leer_filas_template(
        destino
    )

    assert filas[0][1] == Decimal(
        "300.5000"
    )

    assert filas[1][1] == Decimal(
        "200.2500"
    )

    assert sum(
        (
            valor
            for _, valor
            in filas
        ),
        Decimal("0"),
    ) == Decimal(
        "1735.0000"
    )


def test_snapshot_mantiene_semantica_miles_actual(
    tmp_path,
):
    origen = _crear_template(
        tmp_path
        / "miles.xlsx",
        [
            "1,000",
            "200",
            "100",
            "200",
            "100",
            "135",
        ],
    )

    destino = (
        tmp_path
        / "snapshot.xlsx"
    )

    crear_snapshot_template_estable(
        origen,
        destino,
    )

    celdas = _leer_celdas_b(
        destino
    )

    assert celdas[0][0] == 1000
    assert celdas[0][1] == "n"

    filas = _leer_filas_template(
        destino
    )

    assert filas[0][1] == Decimal(
        "1000.0000"
    )


def test_snapshot_porcentaje_conserva_nombre_y_regla(
    tmp_path,
):
    origen = _crear_template(
        tmp_path
        / "plantilla_porcent_prueba.xlsx",
        [
            "20",
            "20",
            "10",
            "20",
            "20",
            "10",
        ],
    )

    nombre_snapshot = (
        _nombre_snapshot_template(
            origen
        )
    )

    assert (
        "_porcent_"
        in nombre_snapshot
    )

    destino = (
        tmp_path
        / nombre_snapshot
    )

    crear_snapshot_template_estable(
        origen,
        destino,
    )

    assert (
        _template_es_porcentaje(
            destino
        )
        is True
    )

    filas = _leer_filas_template(
        destino
    )

    assert sum(
        (
            valor
            for _, valor
            in filas
        ),
        Decimal("0"),
    ) == Decimal(
        "100.0000"
    )

    celdas = _leer_celdas_b(
        destino
    )

    assert all(
        data_type == "n"
        for _, data_type, _
        in celdas
    )


def test_snapshot_numerico_no_necesita_reescritura(
    tmp_path,
):
    origen = _crear_template(
        tmp_path
        / "numerico.xlsx",
        [
            300,
            200,
            100,
            500,
            400,
            235,
        ],
    )

    destino = (
        tmp_path
        / "snapshot.xlsx"
    )

    metadata = crear_snapshot_template_estable(
        origen,
        destino,
    )

    assert (
        metadata[
            "normalizacion_quantity_or_percent"
        ][
            "modificado"
        ]
        is False
    )

    assert (
        metadata[
            "normalizacion_quantity_or_percent"
        ][
            "normalizadas"
        ]
        == 0
    )

    assert (
        metadata[
            "sha256"
        ]
        == metadata[
            "sha256_origen"
        ]
    )



def test_snapshot_template_opaco_conserva_compatibilidad(
    tmp_path,
):
    """
    Compatibilidad con el contrato histórico de snapshot_manager:
    algunos tests usan un archivo opaco con extensión .xlsx
    únicamente para probar deduplicación/rutas.

    No debe intentarse abrir ese fixture con OpenPyXL.
    """
    origen = (
        tmp_path
        / "fixture_opaco.xlsx"
    )

    origen.write_bytes(
        b"TEMPLATE"
    )

    destino = (
        tmp_path
        / "snapshot_opaco.xlsx"
    )

    metadata = (
        crear_snapshot_template_estable(
            origen,
            destino,
        )
    )

    assert (
        destino.read_bytes()
        == b"TEMPLATE"
    )

    normalizacion = metadata[
        "normalizacion_quantity_or_percent"
    ]

    assert (
        normalizacion[
            "modificado"
        ]
        is False
    )

    assert (
        normalizacion[
            "normalizadas"
        ]
        == 0
    )

    assert (
        normalizacion[
            "omitido_no_excel"
        ]
        is True
    )

    assert (
        metadata[
            "sha256"
        ]
        == sha256_archivo(
            destino
        )
    )
