from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from src.excel import excel_transaction
from src.excel.concurrency_guard import (
    capturar_versiones_ordenes,
)
from src.excel.excel_transaction import (
    ExcelArchivoOcupadoError,
    guardar_workbook_atomico,
)
from src.excel.pending_sync import (
    aplicar_actualizaciones_pendientes_a_snapshot,
    encolar_actualizacion_excel,
    obtener_actualizaciones_fallidas,
    obtener_actualizaciones_pendientes,
    persistir_o_encolar_resultado,
    registrar_actualizacion_fallida_excel,
    sincronizar_actualizaciones_pendientes,
)
from src.excel.retry_policy import (
    clasificar_error_persistencia,
)


def _crear_excel(
    ruta: Path,
    *,
    estado: int = 0,
) -> Path:
    wb = Workbook()

    try:
        ws = wb.active
        ws.title = "Ordenes"

        ws.append(
            [
                "ID_ORDEN",
                "ESTADO_RPA",
                "RESUMEN",
                "DESCRIPCION",
            ]
        )

        ws.append(
            [
                "72",
                estado,
                None,
                "OCT-Small Cells GC",
            ]
        )

        adjuntos = wb.create_sheet(
            "Adjuntos"
        )

        adjuntos.append(
            [
                "ID_ORDEN",
                "ARCHIVO",
            ]
        )

        adjuntos.append(
            [
                "72",
                "evidencia.pdf",
            ]
        )

        wb.save(
            ruta
        )

    finally:
        wb.close()

    return ruta


def _error_winerror5() -> PermissionError:
    error = PermissionError(
        13,
        "Acceso denegado simulado",
    )

    error.winerror = 5

    return error


def test_replace_winerror5_es_bloqueo_recuperable(
    tmp_path,
    monkeypatch,
):
    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    monkeypatch.setattr(
        excel_transaction,
        "exigir_excel_disponible",
        lambda *_args, **_kwargs: None,
    )

    def bloquear(
        *_args,
        **_kwargs,
    ):
        raise _error_winerror5()

    monkeypatch.setattr(
        excel_transaction.os,
        "replace",
        bloquear,
    )

    wb = load_workbook(
        excel
    )

    try:
        with pytest.raises(
            ExcelArchivoOcupadoError
        ) as capturado:
            guardar_workbook_atomico(
                wb,
                excel,
            )

        assert (
            capturado.value.winerror
            == 5
        )

    finally:
        wb.close()


def test_retry_policy_winerror5_crudo_no_basta_para_asumir_lock():
    clasificacion = (
        clasificar_error_persistencia(
            _error_winerror5()
        )
    )

    assert clasificacion.recuperable is False

    assert (
        clasificacion.codigo
        == "PERMISO_DENEGADO"
    )

    assert clasificacion.winerror == 5


def test_winerror5_persistente_se_encola_no_se_aisla(
    tmp_path,
    monkeypatch,
):
    import src.excel.state_manager as state_manager

    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path / "journal"
        ),
    )

    monkeypatch.setattr(
        "src.excel.retry_policy.time.sleep",
        lambda *_args, **_kwargs: None,
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    llamadas = {
        "total": 0,
    }

    def bloquear(
        *_args,
        **_kwargs,
    ):
        llamadas[
            "total"
        ] += 1

        raise ExcelArchivoOcupadoError(
            excel,
            (
                "OneDrive bloqueó temporalmente "
                "el reemplazo final."
            ),
            5,
        )

    monkeypatch.setattr(
        state_manager,
        "actualizar_resultado_orden",
        bloquear,
    )

    resultado = (
        persistir_o_encolar_resultado(
            excel,
            "72",
            estado_rpa=2,
            resumen=None,
        )
    )

    assert resultado.encolado is True
    assert resultado.recuperable is True
    assert resultado.intentos == 3
    assert llamadas["total"] == 3

    pendientes = (
        obtener_actualizaciones_pendientes(
            excel
        )
    )

    fallidas = (
        obtener_actualizaciones_fallidas(
            excel
        )
    )

    assert len(pendientes) == 1
    assert pendientes[0]["id_orden"] == "72"
    assert fallidas == []


def test_failed_update_legacy_winerror5_se_recupera(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path / "journal"
        ),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx",
        estado=0,
    )

    version = (
        capturar_versiones_ordenes(
            excel,
            ["72"],
        )["72"]
    )

    registrar_actualizacion_fallida_excel(
        excel,
        "72",
        estado_rpa=2,
        resumen=None,
        codigo_error="EXCEL_ERROR_ACCESO",
        detalle_error=(
            "No se pudo reemplazar DATA.xlsx. "
            "Detalle: [WinError 5] Acceso denegado"
        ),
        tipo_error="ExcelAccesoError",
        intentos_realizados=1,
        version_esperada=version,
    )

    assert len(
        obtener_actualizaciones_fallidas(
            excel
        )
    ) == 1

    # Reproduce el caso real:
    # el Excel ya terminÃ³ mostrando estado 2.
    wb = load_workbook(
        excel
    )

    try:
        wb[
            "Ordenes"
        ][
            "B2"
        ] = 2

        wb.save(
            excel
        )

    finally:
        wb.close()

    resultado = (
        sincronizar_actualizaciones_pendientes(
            excel
        )
    )

    assert resultado["aplicadas"] == 1

    assert (
        obtener_actualizaciones_pendientes(
            excel
        )
        == []
    )

    assert (
        obtener_actualizaciones_fallidas(
            excel
        )
        == []
    )


def test_failed_update_logico_sigue_requiriendo_revision(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path / "journal"
        ),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    registrar_actualizacion_fallida_excel(
        excel,
        "72",
        estado_rpa=2,
        resumen=None,
        codigo_error="DATOS_EXCEL_INVALIDOS",
        detalle_error="Cambio humano",
        tipo_error="ValueError",
        intentos_realizados=1,
    )

    resultado = (
        sincronizar_actualizaciones_pendientes(
            excel
        )
    )

    assert resultado["aplicadas"] == 0

    assert (
        obtener_actualizaciones_pendientes(
            excel
        )
        == []
    )

    fallidas = (
        obtener_actualizaciones_fallidas(
            excel
        )
    )

    assert len(fallidas) == 1
    assert (
        fallidas[0]["sync_error_code"]
        == "DATOS_EXCEL_INVALIDOS"
    )


def test_overlay_snapshot_reintenta_winerror5(
    tmp_path,
    monkeypatch,
):
    import src.excel.state_manager as state_manager

    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path / "journal"
        ),
    )

    monkeypatch.setattr(
        "src.excel.retry_policy.time.sleep",
        lambda *_args, **_kwargs: None,
    )

    original = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    snapshot = (
        tmp_path
        / "entrada_original.xlsx"
    )

    snapshot.write_bytes(
        original.read_bytes()
    )

    encolar_actualizacion_excel(
        original,
        "72",
        estado_rpa=2,
        resumen=None,
    )

    actualizar_real = (
        state_manager.actualizar_resultado_orden
    )

    llamadas = {
        "total": 0,
    }

    def actualizar_transitorio(
        *args,
        **kwargs,
    ):
        llamadas[
            "total"
        ] += 1

        if llamadas[
            "total"
        ] < 3:
            raise ExcelArchivoOcupadoError(
                snapshot,
                "OneDrive simulado",
                5,
            )

        return actualizar_real(
            *args,
            **kwargs,
        )

    monkeypatch.setattr(
        state_manager,
        "actualizar_resultado_orden",
        actualizar_transitorio,
    )

    resultado = (
        aplicar_actualizaciones_pendientes_a_snapshot(
            snapshot,
            original,
        )
    )

    assert resultado["aplicadas"] == 1
    assert resultado["fallidas"] == 0
    assert llamadas["total"] == 3

    wb = load_workbook(
        snapshot,
        data_only=True,
    )

    try:
        assert (
            wb["Ordenes"]["B2"].value
            == 2
        )

    finally:
        wb.close()
