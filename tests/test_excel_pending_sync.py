from __future__ import annotations

from openpyxl import Workbook, load_workbook

from src.excel.excel_transaction import (
    ExcelArchivoOcupadoError,
)

from src.excel.pending_sync import (
    EstadoPersistenciaExcel,
    aplicar_actualizaciones_pendientes_a_snapshot,
    encolar_actualizacion_excel,
    obtener_actualizaciones_fallidas,
    obtener_actualizaciones_pendientes,
    persistir_o_encolar_resultado,
    ruta_journal_excel,
    sincronizar_actualizaciones_pendientes,
)


def _crear_excel(ruta):
    wb = Workbook()

    try:
        ws = wb.active
        ws.title = "Ordenes"

        ws.append(
            [
                "ID_ORDEN",
                "ESTADO_RPA",
                "RESUMEN",
            ]
        )

        ws.append(
            [
                "1001",
                0,
                "",
            ]
        )

        ws.append(
            [
                "1002",
                0,
                "",
            ]
        )

        wb.save(ruta)

    finally:
        wb.close()

    return ruta


def _leer_fila(ruta, id_orden):
    wb = load_workbook(
        ruta,
        data_only=True,
    )

    try:
        ws = wb["Ordenes"]

        for fila in range(
            2,
            ws.max_row + 1,
        ):
            if (
                str(
                    ws.cell(
                        fila,
                        1,
                    ).value
                )
                == str(id_orden)
            ):
                return {
                    "estado": ws.cell(
                        fila,
                        2,
                    ).value,
                    "resumen": ws.cell(
                        fila,
                        3,
                    ).value,
                }

    finally:
        wb.close()

    raise AssertionError(
        f"No se encontró ID {id_orden}"
    )


def test_journal_persiste_actualizacion(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path / "journal"
        ),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    journal = encolar_actualizacion_excel(
        excel,
        "1001",
        estado_rpa=1,
        resumen="OK CBN",
        detalle_error="Excel ocupado",
    )

    assert journal.exists()

    pendientes = (
        obtener_actualizaciones_pendientes(
            excel
        )
    )

    assert len(
        pendientes
    ) == 1

    assert (
        pendientes[0]["id_orden"]
        == "1001"
    )

    assert (
        pendientes[0]["estado_rpa"]
        == 1
    )

    assert (
        pendientes[0]["resumen"]
        == "OK CBN"
    )


def test_journal_deduplica_por_id_y_conserva_ultimo_valor(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path / "journal"
        ),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    encolar_actualizacion_excel(
        excel,
        "1001",
        estado_rpa=2,
        resumen=None,
    )

    encolar_actualizacion_excel(
        excel,
        "1001",
        estado_rpa=1,
        resumen="FINAL",
    )

    pendientes = (
        obtener_actualizaciones_pendientes(
            excel
        )
    )

    assert len(
        pendientes
    ) == 1

    assert (
        pendientes[0]["estado_rpa"]
        == 1
    )

    assert (
        pendientes[0]["resumen"]
        == "FINAL"
    )


def test_persistir_directamente_no_crea_journal(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path / "journal"
        ),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    resultado = (
        persistir_o_encolar_resultado(
            excel,
            "1001",
            estado_rpa=1,
            resumen="DIRECTO",
        )
    )

    assert (
        resultado.estado
        == EstadoPersistenciaExcel.APLICADO
    )

    assert not ruta_journal_excel(
        excel
    ).exists()

    fila = _leer_fila(
        excel,
        "1001",
    )

    assert fila["estado"] == 1
    assert fila["resumen"] == "DIRECTO"


def test_persistir_encola_si_excel_falla(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path / "journal"
        ),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    def fallar(*args, **kwargs):
        raise ExcelArchivoOcupadoError(
            excel,
            "archivo ocupado",
            32,
        )

    monkeypatch.setattr(
        "src.excel.state_manager.actualizar_resultado_orden",
        fallar,
    )

    resultado = (
        persistir_o_encolar_resultado(
            excel,
            "1001",
            estado_rpa=1,
            resumen="CBN OK",
        )
    )

    assert (
        resultado.estado
        == EstadoPersistenciaExcel.ENCOLADO
    )

    pendientes = (
        obtener_actualizaciones_pendientes(
            excel
        )
    )

    assert len(
        pendientes
    ) == 1

    assert (
        pendientes[0]["id_orden"]
        == "1001"
    )


def test_sincronizar_aplica_y_elimina_journal(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path / "journal"
        ),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    encolar_actualizacion_excel(
        excel,
        "1001",
        estado_rpa=1,
        resumen="SINCRONIZADO",
    )

    resultado = (
        sincronizar_actualizaciones_pendientes(
            excel
        )
    )

    assert resultado["total"] == 1
    assert resultado["aplicadas"] == 1
    assert resultado["restantes"] == 0

    fila = _leer_fila(
        excel,
        "1001",
    )

    assert fila["estado"] == 1

    assert (
        fila["resumen"]
        == "SINCRONIZADO"
    )

    assert not ruta_journal_excel(
        excel
    ).exists()


def test_overlay_snapshot_evitar_reprocesamiento(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path / "journal"
        ),
    )

    original = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    snapshot = (
        tmp_path
        / "entrada_original.xlsx"
    )

    snapshot.write_bytes(
        original.read_bytes()
    )

    encolar_actualizacion_excel(
        original,
        "1001",
        estado_rpa=1,
        resumen="PORTAL YA OK",
    )

    resultado = (
        aplicar_actualizaciones_pendientes_a_snapshot(
            snapshot,
            original,
        )
    )

    assert resultado["aplicadas"] == 1

    fila_original = _leer_fila(
        original,
        "1001",
    )

    fila_snapshot = _leer_fila(
        snapshot,
        "1001",
    )

    # El DATA original continúa sin tocarse.
    assert fila_original["estado"] == 0

    # El snapshot sí refleja la realidad del portal.
    assert fila_snapshot["estado"] == 1

    assert (
        fila_snapshot["resumen"]
        == "PORTAL YA OK"
    )

    # El journal sigue existiendo porque todavía
    # no se sincronizó con el archivo original.
    assert ruta_journal_excel(
        original
    ).exists()


def test_persistir_error_logico_no_se_encola(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path / "journal"
        ),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    llamadas = {
        "total": 0,
    }

    def fallar(*args, **kwargs):
        llamadas["total"] += 1

        raise ValueError(
            "No se encontró ID_ORDEN=999"
        )

    monkeypatch.setattr(
        "src.excel.state_manager.actualizar_resultado_orden",
        fallar,
    )

    resultado = (
        persistir_o_encolar_resultado(
            excel,
            "999",
            estado_rpa=1,
            resumen="NO DEBE ENCOLARSE",
        )
    )

    assert (
        resultado.estado
        == EstadoPersistenciaExcel.NO_PERSISTIDO
    )

    assert resultado.recuperable is False

    assert (
        resultado.codigo_error
        == "DATOS_EXCEL_INVALIDOS"
    )

    assert resultado.intentos == 1
    assert llamadas["total"] == 1

    assert not ruta_journal_excel(
        excel
    ).exists()


def test_persistir_bloqueo_reintenta_y_encola(
    tmp_path,
    monkeypatch,
):
    from src.excel.excel_transaction import (
        ExcelArchivoOcupadoError,
    )

    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path / "journal"
        ),
    )

    monkeypatch.setattr(
        "src.excel.retry_policy.time.sleep",
        lambda _: None,
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    llamadas = {
        "total": 0,
    }

    def ocupado(*args, **kwargs):
        llamadas["total"] += 1

        raise ExcelArchivoOcupadoError(
            excel,
            "Excel abierto",
            32,
        )

    monkeypatch.setattr(
        "src.excel.state_manager.actualizar_resultado_orden",
        ocupado,
    )

    resultado = (
        persistir_o_encolar_resultado(
            excel,
            "1001",
            estado_rpa=1,
            resumen="CBN OK",
        )
    )

    assert (
        resultado.estado
        == EstadoPersistenciaExcel.ENCOLADO
    )

    assert resultado.recuperable is True
    assert resultado.intentos == 3

    assert (
        resultado.codigo_error
        == "EXCEL_OCUPADO"
    )

    assert llamadas["total"] == 3

    assert ruta_journal_excel(
        excel
    ).exists()


def test_sync_bloqueo_global_reintenta_solo_primer_id(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path / "journal"
        ),
    )

    monkeypatch.setattr(
        "src.excel.retry_policy.time.sleep",
        lambda _: None,
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    encolar_actualizacion_excel(
        excel,
        "1001",
        estado_rpa=1,
        resumen="UNO",
    )

    encolar_actualizacion_excel(
        excel,
        "1002",
        estado_rpa=1,
        resumen="DOS",
    )

    llamadas = {
        "total": 0,
    }

    def ocupado(*args, **kwargs):
        llamadas["total"] += 1

        raise ExcelArchivoOcupadoError(
            excel,
            "Excel abierto",
            32,
        )

    monkeypatch.setattr(
        "src.excel.state_manager.actualizar_resultado_orden",
        ocupado,
    )

    resultado = (
        sincronizar_actualizaciones_pendientes(
            excel
        )
    )

    # Tres retries únicamente para el primer ID.
    assert llamadas["total"] == 3

    assert resultado["total"] == 2
    assert resultado["aplicadas"] == 0
    assert resultado["recuperables"] == 1
    assert resultado["aisladas"] == 0
    assert resultado["restantes"] == 2

    pendientes = (
        obtener_actualizaciones_pendientes(
            excel
        )
    )

    assert len(pendientes) == 2

    por_id = {
        item["id_orden"]: item
        for item in pendientes
    }

    assert (
        por_id["1001"]["intentos"]
        == 3
    )

    # El segundo no se intentó innecesariamente.
    assert (
        por_id["1002"]["intentos"]
        == 0
    )


def test_sync_error_logico_se_aisla_y_no_se_reintenta(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path / "journal"
        ),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    encolar_actualizacion_excel(
        excel,
        "1001",
        estado_rpa=1,
        resumen="PORTAL OK",
    )

    llamadas = {
        "total": 0,
    }

    def error_logico(*args, **kwargs):
        llamadas["total"] += 1

        raise ValueError(
            "No existe ID_ORDEN=1001"
        )

    monkeypatch.setattr(
        "src.excel.state_manager.actualizar_resultado_orden",
        error_logico,
    )

    primero = (
        sincronizar_actualizaciones_pendientes(
            excel
        )
    )

    assert llamadas["total"] == 1

    assert primero["total"] == 1
    assert primero["recuperables"] == 0
    assert primero["aisladas"] == 1
    assert primero["restantes"] == 0

    assert (
        obtener_actualizaciones_pendientes(
            excel
        )
        == []
    )

    fallidas = (
        obtener_actualizaciones_fallidas(
            excel
        )
    )

    assert len(fallidas) == 1

    assert (
        fallidas[0]["id_orden"]
        == "1001"
    )

    assert (
        fallidas[0]["sync_error_code"]
        == "DATOS_EXCEL_INVALIDOS"
    )

    assert (
        fallidas[0]["sync_recoverable"]
        is False
    )

    # La evidencia debe seguir físicamente guardada.
    assert ruta_journal_excel(
        excel
    ).exists()

    # Segunda sincronización: ya no se intenta.
    segundo = (
        sincronizar_actualizaciones_pendientes(
            excel
        )
    )

    assert segundo["total"] == 0

    assert llamadas["total"] == 1


def test_nueva_actualizacion_reactiva_id_aislado(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path / "journal"
        ),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    encolar_actualizacion_excel(
        excel,
        "1001",
        estado_rpa=1,
        resumen="VERSION 1",
    )

    def fallar(*args, **kwargs):
        raise ValueError(
            "Fallo estructural"
        )

    monkeypatch.setattr(
        "src.excel.state_manager.actualizar_resultado_orden",
        fallar,
    )

    sincronizar_actualizaciones_pendientes(
        excel
    )

    assert len(
        obtener_actualizaciones_fallidas(
            excel
        )
    ) == 1

    # Llega una versión más nueva para el mismo ID.
    encolar_actualizacion_excel(
        excel,
        "1001",
        estado_rpa=1,
        resumen="VERSION 2",
    )

    pendientes = (
        obtener_actualizaciones_pendientes(
            excel
        )
    )

    fallidas = (
        obtener_actualizaciones_fallidas(
            excel
        )
    )

    assert len(pendientes) == 1
    assert pendientes[0]["resumen"] == "VERSION 2"

    assert fallidas == []


def test_persistencia_inicial_registra_retries_en_journal(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path / "journal"
        ),
    )

    monkeypatch.setattr(
        "src.excel.retry_policy.time.sleep",
        lambda _: None,
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    def ocupado(*args, **kwargs):
        raise ExcelArchivoOcupadoError(
            excel,
            "Excel abierto",
            32,
        )

    monkeypatch.setattr(
        "src.excel.state_manager.actualizar_resultado_orden",
        ocupado,
    )

    resultado = (
        persistir_o_encolar_resultado(
            excel,
            "1001",
            estado_rpa=1,
            resumen="CBN OK",
        )
    )

    assert resultado.intentos == 3

    pendientes = (
        obtener_actualizaciones_pendientes(
            excel
        )
    )

    assert len(pendientes) == 1

    # Ahora el journal refleja los retries reales.
    assert (
        pendientes[0]["intentos"]
        == 3
    )
