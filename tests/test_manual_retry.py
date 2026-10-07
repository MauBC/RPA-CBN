from __future__ import annotations

from openpyxl import Workbook, load_workbook
import pytest

from src.excel.pending_sync import (
    ReintentoManualNoPermitidoError,
    encolar_actualizacion_excel,
    preparar_reintento_manual,
    registrar_actualizacion_fallida_excel,
    registrar_orden_inflight,
)


def _crear_excel(
    ruta,
    estados,
):
    wb = Workbook()

    try:
        ws = wb.active
        ws.title = "Ordenes"

        ws.append(
            [
                "ID_ORDEN",
                "ESTADO_RPA",
            ]
        )

        for id_orden, estado in estados:
            ws.append(
                [
                    id_orden,
                    estado,
                ]
            )

        wb.save(
            ruta
        )

    finally:
        wb.close()

    return ruta


def _estados_id(
    ruta,
    id_orden,
):
    wb = load_workbook(
        ruta,
        data_only=True,
    )

    try:
        ws = wb[
            "Ordenes"
        ]

        return [
            ws.cell(
                fila,
                2,
            ).value
            for fila in range(
                2,
                ws.max_row + 1,
            )
            if str(
                ws.cell(
                    fila,
                    1,
                ).value
            ) == str(
                id_orden
            )
        ]

    finally:
        wb.close()


def test_reintento_manual_cambia_todas_las_posiciones_2_a_0(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path
            / "journal"
        ),
    )

    excel = _crear_excel(
        tmp_path
        / "DATA.xlsx",
        [
            ("1001", 2),
            ("1001", 2),
            ("1002", 0),
        ],
    )

    filas = preparar_reintento_manual(
        excel,
        "1001",
    )

    assert filas == 2

    assert _estados_id(
        excel,
        "1001",
    ) == [
        0,
        0,
    ]

    assert _estados_id(
        excel,
        "1002",
    ) == [
        0,
    ]


def test_reintento_manual_rechaza_orden_que_no_esta_en_error(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path
            / "journal"
        ),
    )

    excel = _crear_excel(
        tmp_path
        / "DATA.xlsx",
        [
            ("1001", 0),
            ("1001", 0),
        ],
    )

    with pytest.raises(
        ValueError,
        match="ESTADO_RPA=2",
    ):
        preparar_reintento_manual(
            excel,
            "1001",
        )

    assert _estados_id(
        excel,
        "1001",
    ) == [
        0,
        0,
    ]


@pytest.mark.parametrize(
    "tipo_bloqueo",
    [
        "updates",
        "failed_updates",
        "inflight",
    ],
)
def test_reintento_manual_bloquea_evidencia_persistente(
    tmp_path,
    monkeypatch,
    tipo_bloqueo,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path
            / "journal"
        ),
    )

    excel = _crear_excel(
        tmp_path
        / "DATA.xlsx",
        [
            ("1001", 2),
        ],
    )

    if tipo_bloqueo == "updates":
        encolar_actualizacion_excel(
            excel,
            "1001",
            estado_rpa=2,
            resumen=None,
            detalle_error="prueba",
        )

    elif tipo_bloqueo == "failed_updates":
        registrar_actualizacion_fallida_excel(
            excel,
            "1001",
            estado_rpa=2,
            resumen=None,
            codigo_error="TEST",
            detalle_error="prueba",
            tipo_error="RuntimeError",
        )

    else:
        registrar_orden_inflight(
            excel,
            "1001",
        )

    with pytest.raises(
        ReintentoManualNoPermitidoError,
        match="No es seguro reintentar",
    ):
        preparar_reintento_manual(
            excel,
            "1001",
        )

    assert _estados_id(
        excel,
        "1001",
    ) == [
        2,
    ]
