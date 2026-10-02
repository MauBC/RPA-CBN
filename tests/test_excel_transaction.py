from __future__ import annotations

from openpyxl import Workbook, load_workbook
import pytest

from src.excel.excel_transaction import (
    ExcelArchivoOcupadoError,
    guardar_workbook_atomico,
)


def _crear_excel(ruta):
    wb = Workbook()

    try:
        ws = wb.active
        ws.title = "Ordenes"
        ws["A1"] = "VALOR"
        ws["A2"] = "ORIGINAL"

        wb.save(ruta)

    finally:
        wb.close()


def test_guardado_atomico_actualiza_excel(tmp_path):
    ruta = tmp_path / "DATA.xlsx"
    _crear_excel(ruta)

    wb = load_workbook(ruta)

    try:
        wb["Ordenes"]["A2"] = "NUEVO"

        guardar_workbook_atomico(
            wb,
            ruta,
        )
    finally:
        wb.close()

    verificacion = load_workbook(
        ruta,
        data_only=True,
    )

    try:
        assert (
            verificacion["Ordenes"]["A2"].value
            == "NUEVO"
        )
    finally:
        verificacion.close()


def test_guardado_atomico_no_deja_temporales(tmp_path):
    ruta = tmp_path / "DATA.xlsx"
    _crear_excel(ruta)

    wb = load_workbook(ruta)

    try:
        wb["Ordenes"]["A2"] = "NUEVO"

        guardar_workbook_atomico(
            wb,
            ruta,
        )
    finally:
        wb.close()

    temporales = list(
        tmp_path.glob(
            ".DATA.rpa-*"
        )
    )

    assert temporales == []


def test_guardado_atomico_preserva_original_si_archivo_esta_ocupado(
    tmp_path,
    monkeypatch,
):
    ruta = tmp_path / "DATA.xlsx"
    _crear_excel(ruta)

    contenido_original = (
        ruta.read_bytes()
    )

    wb = load_workbook(ruta)

    try:
        wb["Ordenes"]["A2"] = "NO DEBE GUARDARSE"

        def simular_ocupado(_ruta):
            raise ExcelArchivoOcupadoError(
                _ruta,
                "Simulación de archivo ocupado.",
                32,
            )

        monkeypatch.setattr(
            "src.excel.excel_transaction.exigir_excel_disponible",
            simular_ocupado,
        )

        with pytest.raises(
            ExcelArchivoOcupadoError
        ):
            guardar_workbook_atomico(
                wb,
                ruta,
            )

    finally:
        wb.close()

    assert (
        ruta.read_bytes()
        == contenido_original
    )


def test_guardado_atomico_el_temporal_es_excel_valido(
    tmp_path,
):
    ruta = tmp_path / "DATA.xlsx"
    _crear_excel(ruta)

    wb = load_workbook(ruta)

    try:
        wb["Ordenes"]["A2"] = "VALIDO"

        guardar_workbook_atomico(
            wb,
            ruta,
        )
    finally:
        wb.close()

    wb_final = load_workbook(
        ruta,
        read_only=True,
        data_only=True,
    )

    try:
        assert (
            wb_final["Ordenes"]["A2"].value
            == "VALIDO"
        )
    finally:
        wb_final.close()


def test_guardado_cas_aplica_si_archivo_no_cambio(
    tmp_path,
):
    from src.excel.excel_transaction import (
        sha256_archivo_excel,
    )

    ruta = tmp_path / "DATA_CAS.xlsx"
    _crear_excel(ruta)

    sha = sha256_archivo_excel(
        ruta
    )

    wb = load_workbook(
        ruta
    )

    try:
        wb["Ordenes"]["A2"] = "RPA"

        guardar_workbook_atomico(
            wb,
            ruta,
            sha256_esperado=sha,
        )

    finally:
        wb.close()

    verificacion = load_workbook(
        ruta,
        data_only=True,
    )

    try:
        assert (
            verificacion["Ordenes"]["A2"].value
            == "RPA"
        )
    finally:
        verificacion.close()


def test_guardado_cas_no_pisa_cambio_externo(
    tmp_path,
):
    from src.excel.excel_transaction import (
        ExcelArchivoCambioConcurrenteError,
        sha256_archivo_excel,
    )

    ruta = tmp_path / "DATA_CAS.xlsx"
    _crear_excel(ruta)

    sha = sha256_archivo_excel(
        ruta
    )

    wb_rpa = load_workbook(
        ruta
    )

    try:
        wb_rpa["Ordenes"]["A2"] = (
            "CAMBIO RPA"
        )

        # Otro actor modifica DATA.xlsx mientras
        # nuestro workbook viejo continúa en memoria.
        wb_humano = load_workbook(
            ruta
        )

        try:
            wb_humano["Ordenes"]["A2"] = (
                "CAMBIO HUMANO"
            )

            wb_humano.save(
                ruta
            )

        finally:
            wb_humano.close()

        with pytest.raises(
            ExcelArchivoCambioConcurrenteError,
        ):
            guardar_workbook_atomico(
                wb_rpa,
                ruta,
                sha256_esperado=sha,
            )

    finally:
        wb_rpa.close()

    verificacion = load_workbook(
        ruta,
        data_only=True,
    )

    try:
        # El cambio externo debe sobrevivir.
        assert (
            verificacion["Ordenes"]["A2"].value
            == "CAMBIO HUMANO"
        )

    finally:
        verificacion.close()
