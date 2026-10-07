from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from src.excel.pending_sync import (
    CambioEstadoIgnoradaNoPermitidoError,
    encolar_actualizacion_excel,
    ignorar_ordenes_manual,
    reabrir_ordenes_manual,
)
from src.excel.state_manager import (
    actualizar_estado_orden,
    ignorar_ordenes_fallidas,
    inspeccionar_ids_ordenes_fallidas,
    inspeccionar_ids_ordenes_ignoradas,
    inspeccionar_pendientes,
    reabrir_ordenes_ignoradas,
)


def _crear_excel(
    ruta: Path,
    filas,
) -> Path:
    wb = Workbook()

    try:
        ws = wb.active
        ws.title = "Ordenes"

        ws.append(
            [
                "ID_ORDEN",
                "ESTADO_RPA",
            ]
        )

        for id_orden, estado in filas:
            ws.append(
                [
                    id_orden,
                    estado,
                ]
            )

        wb.save(
            ruta
        )

    finally:
        wb.close()

    return ruta


def _estados(
    ruta: Path,
    id_orden: str,
) -> list[int]:
    wb = load_workbook(
        ruta,
        data_only=True,
    )

    try:
        ws = wb[
            "Ordenes"
        ]

        return [
            ws.cell(
                row=fila,
                column=2,
            ).value
            for fila in range(
                2,
                ws.max_row + 1,
            )
            if str(
                ws.cell(
                    row=fila,
                    column=1,
                ).value
            ) == str(
                id_orden
            )
        ]

    finally:
        wb.close()


def test_cp15_ignora_multiples_ordenes_en_una_transaccion(
    tmp_path,
):
    excel = _crear_excel(
        tmp_path / "DATA.xlsx",
        [
            ("1001", 2),
            ("1001", 2),
            ("1002", 2),
            ("1003", 0),
        ],
    )

    resultado = ignorar_ordenes_fallidas(
        excel,
        [
            "1001",
            "1002",
        ],
    )

    assert resultado == {
        "1001": 2,
        "1002": 1,
    }

    assert _estados(
        excel,
        "1001",
    ) == [
        3,
        3,
    ]

    assert _estados(
        excel,
        "1002",
    ) == [
        3,
    ]

    assert _estados(
        excel,
        "1003",
    ) == [
        0,
    ]


def test_cp15_estado3_no_es_pendiente_ni_error(
    tmp_path,
):
    excel = _crear_excel(
        tmp_path / "DATA.xlsx",
        [
            ("1001", 3),
            ("1002", 0),
            ("1003", 2),
        ],
    )

    pendientes = inspeccionar_pendientes(
        excel
    )

    assert [
        item.id_orden
        for item in pendientes
    ] == [
        "1002"
    ]

    assert (
        inspeccionar_ids_ordenes_fallidas(
            excel
        )
        == [
            "1003"
        ]
    )

    assert (
        inspeccionar_ids_ordenes_ignoradas(
            excel
        )
        == [
            "1001"
        ]
    )


def test_cp15_reabre_ignorada_de_3_a_2(
    tmp_path,
):
    excel = _crear_excel(
        tmp_path / "DATA.xlsx",
        [
            ("1001", 3),
            ("1001", 3),
        ],
    )

    resultado = reabrir_ordenes_ignoradas(
        excel,
        [
            "1001"
        ],
    )

    assert resultado == {
        "1001": 2,
    }

    assert _estados(
        excel,
        "1001",
    ) == [
        2,
        2,
    ]


def test_cp15_no_ignora_orden_que_no_esta_en_2(
    tmp_path,
):
    excel = _crear_excel(
        tmp_path / "DATA.xlsx",
        [
            ("1001", 1),
        ],
    )

    with pytest.raises(
        ValueError,
        match="ESTADO_RPA=2",
    ):
        ignorar_ordenes_fallidas(
            excel,
            [
                "1001"
            ],
        )

    assert _estados(
        excel,
        "1001",
    ) == [
        1
    ]


def test_cp15_journal_bloquea_ignorar(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(
            tmp_path / "journal"
        ),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx",
        [
            ("1001", 2),
        ],
    )

    encolar_actualizacion_excel(
        excel,
        "1001",
        estado_rpa=2,
        resumen=None,
        detalle_error="prueba",
    )

    with pytest.raises(
        CambioEstadoIgnoradaNoPermitidoError
    ):
        ignorar_ordenes_manual(
            excel,
            [
                "1001"
            ],
        )

    assert _estados(
        excel,
        "1001",
    ) == [
        2
    ]


def test_cp15_actualizar_estado_acepta_3(
    tmp_path,
):
    excel = _crear_excel(
        tmp_path / "DATA.xlsx",
        [
            ("1001", 2),
        ],
    )

    actualizar_estado_orden(
        excel,
        "1001",
        3,
    )

    assert _estados(
        excel,
        "1001",
    ) == [
        3
    ]
