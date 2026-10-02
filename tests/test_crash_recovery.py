from __future__ import annotations

from pathlib import Path

from concurrent.futures import ThreadPoolExecutor
import os
import time

from openpyxl import Workbook, load_workbook
import pytest

import src.excel.pending_sync as pending_sync
import src.excel.state_manager as state_manager

from src.excel.concurrency_guard import (
    capturar_versiones_ordenes,
)
from src.excel.pending_sync import (
    encolar_actualizacion_excel,
    obtener_actualizaciones_fallidas,
    obtener_actualizaciones_pendientes,
    sincronizar_actualizaciones_pendientes,
)
from src.excel.state_manager import (
    actualizar_resultado_orden,
)


def _crear_excel(
    ruta: Path,
) -> Path:
    wb = Workbook()

    try:
        ws = wb.active
        ws.title = "Ordenes"

        ws.append(
            [
                "ID_ORDEN",
                "ESTADO_RPA",
                "RESUMEN",
                "NOTA",
            ]
        )

        ws.append(
            [
                "1001",
                0,
                "",
                "",
            ]
        )

        ws.append(
            [
                "1002",
                0,
                "",
                "",
            ]
        )

        wb.save(
            ruta
        )

    finally:
        wb.close()

    return ruta


def test_journal_lock_huerfano_se_recupera(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(tmp_path / "journal"),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    lock = pending_sync._ruta_lock_journal(
        excel
    )

    lock.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    lock.write_text(
        "pid=999999 time=0",
        encoding="utf-8",
    )

    viejo = time.time() - 180

    os.utime(
        lock,
        (
            viejo,
            viejo,
        ),
    )

    encolar_actualizacion_excel(
        excel,
        "1001",
        estado_rpa=1,
        resumen="OK",
    )

    assert not lock.exists()

    pendientes = (
        obtener_actualizaciones_pendientes(
            excel
        )
    )

    assert len(pendientes) == 1


def test_journal_lock_activo_no_se_roba(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(tmp_path / "journal"),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    lock = pending_sync._ruta_lock_journal(
        excel
    )

    lock.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    lock.write_text(
        "pid=123 time=actual",
        encoding="utf-8",
    )

    try:
        with pytest.raises(
            TimeoutError
        ):
            with pending_sync._journal_lock(
                excel,
                timeout_segundos=0.05,
            ):
                pass

        assert lock.exists()

    finally:
        lock.unlink(
            missing_ok=True
        )


def test_encolado_concurrente_no_pierde_actualizaciones(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(tmp_path / "journal"),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    ids = [
        str(2000 + indice)
        for indice
        in range(12)
    ]

    def encolar(
        id_orden: str,
    ):
        return encolar_actualizacion_excel(
            excel,
            id_orden,
            estado_rpa=1,
            resumen=f"RESULTADO {id_orden}",
        )

    with ThreadPoolExecutor(
        max_workers=6
    ) as executor:
        resultados = list(
            executor.map(
                encolar,
                ids,
            )
        )

    assert all(
        ruta.exists()
        for ruta in resultados
    )

    pendientes = (
        obtener_actualizaciones_pendientes(
            excel
        )
    )

    encontrados = {
        item["id_orden"]
        for item in pendientes
    }

    assert encontrados == set(ids)
    assert len(pendientes) == 12


def test_excel_lock_huerfano_se_recupera(
    tmp_path,
):
    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    lock = Path(
        str(excel) + ".lock"
    )

    lock.write_text(
        "pid=999999 time=0",
        encoding="utf-8",
    )

    viejo = time.time() - 700

    os.utime(
        lock,
        (
            viejo,
            viejo,
        ),
    )

    with state_manager._excel_lock(
        excel,
        timeout_segundos=1,
    ):
        assert lock.exists()

    assert not lock.exists()


def test_excel_lock_activo_no_se_roba(
    tmp_path,
):
    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    lock = Path(
        str(excel) + ".lock"
    )

    with state_manager._excel_lock(
        excel,
        timeout_segundos=1,
    ):
        assert lock.exists()

        with pytest.raises(
            TimeoutError
        ):
            with state_manager._excel_lock(
                excel,
                timeout_segundos=0,
            ):
                pass

        # El intento secundario no debe borrar
        # el lock del propietario original.
        assert lock.exists()

    assert not lock.exists()


def test_crash_despues_de_excel_antes_de_borrar_journal_es_idempotente(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(tmp_path / "journal"),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    version = (
        capturar_versiones_ordenes(
            excel,
            ["1001"],
        )["1001"]
    )

    # CBN ya produjo resultado y el journal
    # qued? preparado para sincronizar.
    encolar_actualizacion_excel(
        excel,
        "1001",
        estado_rpa=1,
        resumen="PORTAL OK",
        version_esperada=version,
    )

    # Simulamos:
    #
    # 1. Excel fue actualizado correctamente.
    # 2. El proceso muri? ANTES de eliminar el journal.
    actualizar_resultado_orden(
        excel,
        "1001",
        estado_rpa=1,
        resumen="PORTAL OK",
        version_esperada=version,
    )

    assert len(
        obtener_actualizaciones_pendientes(
            excel
        )
    ) == 1

    # Reinicio del proceso.
    resultado = (
        sincronizar_actualizaciones_pendientes(
            excel
        )
    )

    assert resultado["aplicadas"] == 1
    assert resultado["restantes"] == 0

    assert (
        obtener_actualizaciones_pendientes(
            excel
        )
        == []
    )

    assert (
        obtener_actualizaciones_fallidas(
            excel
        )
        == []
    )

    wb = load_workbook(
        excel,
        data_only=True,
    )

    try:
        ws = wb["Ordenes"]

        assert ws["B2"].value == 1
        assert ws["C2"].value == "PORTAL OK"

    finally:
        wb.close()


def test_replay_crash_preserva_edicion_humana_de_otra_orden(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(tmp_path / "journal"),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    version = (
        capturar_versiones_ordenes(
            excel,
            ["1001"],
        )["1001"]
    )

    encolar_actualizacion_excel(
        excel,
        "1001",
        estado_rpa=1,
        resumen="PORTAL OK",
        version_esperada=version,
    )

    actualizar_resultado_orden(
        excel,
        "1001",
        estado_rpa=1,
        resumen="PORTAL OK",
        version_esperada=version,
    )

    # Despu?s del commit del RPA, pero antes de
    # limpiar el journal, una persona modifica 1002.
    wb = load_workbook(
        excel
    )

    try:
        ws = wb["Ordenes"]

        ws["D3"] = (
            "CAMBIO HUMANO PRESERVADO"
        )

        wb.save(
            excel
        )

    finally:
        wb.close()

    resultado = (
        sincronizar_actualizaciones_pendientes(
            excel
        )
    )

    assert resultado["aplicadas"] == 1
    assert resultado["restantes"] == 0

    wb = load_workbook(
        excel,
        data_only=True,
    )

    try:
        ws = wb["Ordenes"]

        assert ws["B2"].value == 1
        assert ws["C2"].value == "PORTAL OK"

        assert (
            ws["D3"].value
            == "CAMBIO HUMANO PRESERVADO"
        )

    finally:
        wb.close()



def test_journal_lock_reintenta_permissionerror_transitorio_windows(
    tmp_path,
    monkeypatch,
):
    import src.excel.pending_sync as pending_sync

    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(tmp_path / "journal"),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    original_open = (
        pending_sync.os.open
    )

    llamadas = {
        "total": 0,
    }

    def open_con_colision(
        path,
        flags,
        *args,
        **kwargs,
    ):
        if (
            str(path).endswith(".lock")
            and llamadas["total"] == 0
        ):
            llamadas["total"] += 1

            raise PermissionError(
                13,
                "Colision transitoria Windows",
                str(path),
            )

        llamadas["total"] += 1

        return original_open(
            path,
            flags,
            *args,
            **kwargs,
        )

    monkeypatch.setattr(
        pending_sync.os,
        "open",
        open_con_colision,
    )

    with pending_sync._journal_lock(
        excel,
        timeout_segundos=1,
    ):
        pass

    assert llamadas["total"] >= 2

    assert not pending_sync._ruta_lock_journal(
        excel
    ).exists()
