from openpyxl import load_workbook
import pytest

from src.excel.state_manager import (
    actualizar_estado_orden,
    obtener_fila_excel_por_id,
    obtener_pendientes,
    preparar_columna_estado_y_validar_ids,
)

from conftest import crear_excel_estados


def test_crea_columna_estado_si_no_existe(tmp_path):
    ruta = crear_excel_estados(
        tmp_path / "DATA.xlsx",
        [
            ("1001",),
            ("1002",),
        ],
        incluir_estado=False,
    )

    preparar_columna_estado_y_validar_ids(ruta)

    wb = load_workbook(ruta, data_only=True)
    try:
        ws = wb["Ordenes"]

        headers = {
            str(cell.value).strip(): cell.column
            for cell in ws[1]
            if cell.value is not None
        }

        assert "ESTADO_RPA" in headers

        col_estado = headers["ESTADO_RPA"]

        assert ws.cell(2, col_estado).value == 0
        assert ws.cell(3, col_estado).value == 0
    finally:
        wb.close()


def test_actualizar_estado_actualiza_todas_las_posiciones(tmp_path):
    ruta = crear_excel_estados(
        tmp_path / "DATA.xlsx",
        [
            ("1001", 0),
            ("1001", 0),
            ("1002", 0),
        ],
    )

    actualizar_estado_orden(
        ruta_excel=ruta,
        id_orden_buscado="1001",
        estado_rpa=1,
    )

    wb = load_workbook(ruta, data_only=True)
    try:
        ws = wb["Ordenes"]

        assert ws["B2"].value == 1
        assert ws["B3"].value == 1
        assert ws["B4"].value == 0
    finally:
        wb.close()


def test_actualizar_estado_rechaza_estado_invalido(tmp_path):
    ruta = crear_excel_estados(
        tmp_path / "DATA.xlsx",
        [
            ("1001", 0),
        ],
    )

    with pytest.raises(ValueError):
        actualizar_estado_orden(
            ruta_excel=ruta,
            id_orden_buscado="1001",
            estado_rpa=9,
        )


def test_detecta_estados_inconsistentes_para_mismo_id(tmp_path):
    ruta = crear_excel_estados(
        tmp_path / "DATA.xlsx",
        [
            ("1001", 0),
            ("1001", 1),
        ],
    )

    with pytest.raises(
        ValueError,
        match="mismo ID_ORDEN",
    ):
        preparar_columna_estado_y_validar_ids(ruta)


def test_obtener_pendientes_devuelve_un_elemento_por_id(tmp_path):
    ruta = crear_excel_estados(
        tmp_path / "DATA.xlsx",
        [
            ("1001", 0),
            ("1001", 0),
            ("1002", 1),
            ("1003", 0),
            ("1003", 0),
        ],
    )

    pendientes = obtener_pendientes(ruta)

    ids = [item.id_orden for item in pendientes]

    assert ids == ["1001", "1003"]


def test_obtener_fila_excel_por_id(tmp_path):
    ruta = crear_excel_estados(
        tmp_path / "DATA.xlsx",
        [
            ("1001", 0),
            ("1001", 0),
            ("1002", 0),
        ],
    )

    fila = obtener_fila_excel_por_id(
        ruta,
        "1002",
    )

    assert fila == 4


def test_obtener_fila_libera_excel_despues_de_lectura(tmp_path):
    ruta = crear_excel_estados(
        tmp_path / "DATA.xlsx",
        [
            ("1001", 0),
            ("1002", 0),
        ],
    )

    fila = obtener_fila_excel_por_id(
        ruta,
        "1002",
    )

    assert fila == 3

    destino = tmp_path / "DATA_RENAMED.xlsx"
    ruta.rename(destino)

    assert destino.exists()


def test_obtener_pendientes_libera_excel_despues_de_lectura(tmp_path):
    ruta = crear_excel_estados(
        tmp_path / "DATA.xlsx",
        [
            ("1001", 0),
            ("1001", 0),
            ("1002", 1),
        ],
    )

    pendientes = obtener_pendientes(ruta)

    assert [p.id_orden for p in pendientes] == ["1001"]

    destino = tmp_path / "DATA_RENAMED.xlsx"
    ruta.rename(destino)

    assert destino.exists()


def test_actualizar_estado_libera_excel_despues_de_guardar(tmp_path):
    ruta = crear_excel_estados(
        tmp_path / "DATA.xlsx",
        [
            ("1001", 0),
        ],
    )

    actualizar_estado_orden(
        ruta,
        "1001",
        1,
    )

    destino = tmp_path / "DATA_RENAMED.xlsx"
    ruta.rename(destino)

    assert destino.exists()
