from __future__ import annotations

import inspect
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.core import runner
from src.excel.file_access import (
    DiagnosticoAccesoExcel,
    EstadoAccesoExcel,
)
from src.excel.pending_sync import (
    EstadoPersistenciaExcel,
    ResultadoPersistenciaExcel,
    aplicar_actualizaciones_pendientes_a_snapshot,
    cerrar_orden_inflight,
    encolar_actualizacion_excel,
    obtener_estado_journal,
    obtener_ordenes_inflight,
    registrar_orden_inflight,
    ruta_journal_excel,
    sincronizar_actualizaciones_pendientes,
)
from src.excel.sync_status import (
    EstadoSincronizacionExcel,
    obtener_estado_sincronizacion_excel,
)


def _crear_excel(
    ruta: Path,
) -> Path:
    wb = Workbook()

    try:
        ws = wb.active
        ws.title = "Ordenes"

        ws.append(
            [
                "ID_ORDEN",
                "ESTADO_RPA",
                "RESUMEN",
            ]
        )

        ws.append(
            [
                "1001",
                0,
                "",
            ]
        )

        ws.append(
            [
                "1002",
                0,
                "",
            ]
        )

        wb.save(
            ruta
        )

    finally:
        wb.close()

    return ruta


def test_inflight_se_registra_y_se_cierra(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(tmp_path / "journal"),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    token = registrar_orden_inflight(
        excel,
        "1001",
    )

    estado = obtener_estado_journal(
        excel
    )

    assert len(
        estado["inflight"]
    ) == 1

    assert (
        estado["inflight"][0]["id_orden"]
        == "1001"
    )

    assert (
        estado["inflight"][0]["inflight_id"]
        == token
    )

    assert estado["pendientes"] == []
    assert estado["fallidas"] == []

    assert cerrar_orden_inflight(
        excel,
        "1001",
        token,
    )

    assert (
        obtener_ordenes_inflight(
            excel
        )
        == []
    )

    assert not ruta_journal_excel(
        excel
    ).exists()


def test_token_inflight_incorrecto_no_borra_marca(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(tmp_path / "journal"),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    registrar_orden_inflight(
        excel,
        "1001",
    )

    assert (
        cerrar_orden_inflight(
            excel,
            "1001",
            "TOKEN-INCORRECTO",
        )
        is False
    )

    assert len(
        obtener_ordenes_inflight(
            excel
        )
    ) == 1


def test_overlay_inflight_excluye_orden_sin_tocar_original(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(tmp_path / "journal"),
    )

    original = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    snapshot = (
        tmp_path
        / "entrada_original.xlsx"
    )

    snapshot.write_bytes(
        original.read_bytes()
    )

    registrar_orden_inflight(
        original,
        "1001",
    )

    resultado = (
        aplicar_actualizaciones_pendientes_a_snapshot(
            snapshot,
            original,
        )
    )

    assert resultado["inflight"] == 1

    assert (
        resultado["omitidas_inflight"]
        == 1
    )

    assert resultado["fallidas"] == 0

    wb_original = load_workbook(
        original,
        data_only=True,
    )

    wb_snapshot = load_workbook(
        snapshot,
        data_only=True,
    )

    try:
        assert (
            wb_original["Ordenes"]["B2"].value
            == 0
        )

        # Solo el snapshot excluye la orden incierta.
        assert (
            wb_snapshot["Ordenes"]["B2"].value
            == 2
        )

        # 1002 sigue siendo procesable.
        assert (
            wb_snapshot["Ordenes"]["B3"].value
            == 0
        )

    finally:
        wb_original.close()
        wb_snapshot.close()


def test_sync_depura_inflight_si_update_ya_es_durable(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setenv(
        "RPA_CBN_PENDING_DIR",
        str(tmp_path / "journal"),
    )

    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    registrar_orden_inflight(
        excel,
        "1001",
    )

    # Simula:
    # resultado durable en journal,
    # pero crash antes de borrar inflight.
    encolar_actualizacion_excel(
        excel,
        "1001",
        estado_rpa=1,
        resumen="PORTAL OK",
    )

    assert len(
        obtener_ordenes_inflight(
            excel
        )
    ) == 1

    resultado = (
        sincronizar_actualizaciones_pendientes(
            excel
        )
    )

    assert resultado["aplicadas"] == 1
    assert resultado["restantes"] == 0

    estado = obtener_estado_journal(
        excel
    )

    assert estado["inflight"] == []
    assert estado["pendientes"] == []


def test_helper_cierra_inflight_si_excel_queda_aplicado(
    monkeypatch,
):
    llamadas = {
        "cerrar": 0,
    }

    persistencia = (
        ResultadoPersistenciaExcel(
            estado=(
                EstadoPersistenciaExcel.APLICADO
            ),
            id_orden="1001",
            detalle="OK",
            ruta_journal=None,
            intentos=1,
            codigo_error=None,
            recuperable=None,
        )
    )

    monkeypatch.setattr(
        runner,
        "persistir_o_encolar_resultado",
        lambda *args, **kwargs: persistencia,
    )

    def cerrar(
        *args,
        **kwargs,
    ):
        llamadas["cerrar"] += 1
        return True

    monkeypatch.setattr(
        runner,
        "cerrar_orden_inflight",
        cerrar,
    )

    runner._persistir_resultado_excel(
        Path("DATA.xlsx"),
        "1001",
        estado_rpa=1,
        resumen="OK",
        inflight_id="token",
    )

    assert llamadas["cerrar"] == 1


def test_helper_cierra_inflight_si_resultado_queda_en_journal(
    monkeypatch,
):
    llamadas = {
        "cerrar": 0,
    }

    persistencia = (
        ResultadoPersistenciaExcel(
            estado=(
                EstadoPersistenciaExcel.ENCOLADO
            ),
            id_orden="1001",
            detalle="pendiente",
            ruta_journal=Path(
                "journal.json"
            ),
            intentos=3,
            codigo_error="EXCEL_OCUPADO",
            recuperable=True,
        )
    )

    monkeypatch.setattr(
        runner,
        "persistir_o_encolar_resultado",
        lambda *args, **kwargs: persistencia,
    )

    def cerrar(
        *args,
        **kwargs,
    ):
        llamadas["cerrar"] += 1
        return True

    monkeypatch.setattr(
        runner,
        "cerrar_orden_inflight",
        cerrar,
    )

    runner._persistir_resultado_excel(
        Path("DATA.xlsx"),
        "1001",
        estado_rpa=1,
        resumen="OK",
        inflight_id="token",
    )

    assert llamadas["cerrar"] == 1


def test_helper_conserva_inflight_si_no_existe_evidencia_durable(
    monkeypatch,
):
    llamadas = {
        "cerrar": 0,
    }

    persistencia = (
        ResultadoPersistenciaExcel(
            estado=(
                EstadoPersistenciaExcel.NO_PERSISTIDO
            ),
            id_orden="1001",
            detalle="fallo total",
            ruta_journal=None,
            intentos=1,
            codigo_error="ERROR",
            recuperable=False,
        )
    )

    monkeypatch.setattr(
        runner,
        "persistir_o_encolar_resultado",
        lambda *args, **kwargs: persistencia,
    )

    def cerrar(
        *args,
        **kwargs,
    ):
        llamadas["cerrar"] += 1
        return True

    monkeypatch.setattr(
        runner,
        "cerrar_orden_inflight",
        cerrar,
    )

    runner._persistir_resultado_excel(
        Path("DATA.xlsx"),
        "1001",
        estado_rpa=1,
        resumen="OK",
        inflight_id="token",
    )

    assert llamadas["cerrar"] == 0


def test_runner_registra_inflight_antes_del_portal():
    codigo = inspect.getsource(
        runner.ejecutar_rpa
    )

    pos_inflight = codigo.index(
        "registrar_orden_inflight("
    )

    pos_portal = codigo.index(
        "procesar_orden("
    )

    pos_persistencia = codigo.index(
        "_persistir_resultado_excel(",
        pos_portal,
    )

    assert (
        pos_inflight
        < pos_portal
        < pos_persistencia
    )

    assert (
        codigo.count(
            "inflight_id=inflight_id_orden"
        )
        == 3
    )


def test_sync_status_inflight_requiere_revision(
    tmp_path,
    monkeypatch,
):
    excel = _crear_excel(
        tmp_path / "DATA.xlsx"
    )

    diagnostico = (
        DiagnosticoAccesoExcel(
            ruta=excel,
            estado=(
                EstadoAccesoExcel.DISPONIBLE
            ),
            puede_escribir=True,
            lock_office_detectado=False,
            detalle="OK",
            winerror=None,
        )
    )

    monkeypatch.setattr(
        "src.excel.sync_status.diagnosticar_acceso_excel",
        lambda _: diagnostico,
    )

    monkeypatch.setattr(
        "src.excel.sync_status.obtener_estado_journal",
        lambda _: {
            "pendientes": [],
            "fallidas": [],
            "inflight": [
                {
                    "id_orden": "1001",
                }
            ],
        },
    )

    resultado = (
        obtener_estado_sincronizacion_excel(
            excel
        )
    )

    assert (
        resultado.estado
        == EstadoSincronizacionExcel.REQUIERE_REVISION
    )

    assert resultado.requiere_atencion is True
    assert "incierta" in resultado.mensaje
