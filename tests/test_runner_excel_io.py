from pathlib import Path
from types import SimpleNamespace
import inspect

from openpyxl import Workbook

from src.core import runner
import src.excel.concurrency_guard as concurrency_guard
import src.excel.state_manager as state_manager


def test_indexar_filas_pendientes():
    pendientes = [
        SimpleNamespace(
            id_orden="1001",
            fila_excel=2,
        ),
        SimpleNamespace(
            id_orden="1002",
            fila_excel=5,
        ),
    ]

    resultado = (
        runner._indexar_filas_pendientes(
            pendientes
        )
    )

    assert resultado == {
        "1001": 2,
        "1002": 5,
    }


def test_resultado_ok_usa_fila_cacheada_sin_abrir_excel(
    monkeypatch,
):
    def prohibido(*args, **kwargs):
        raise AssertionError(
            "No debe abrir Excel para buscar la fila."
        )

    monkeypatch.setattr(
        runner,
        "obtener_fila_excel_por_id",
        prohibido,
    )

    orden = SimpleNamespace(
        id_orden="1001"
    )

    resultado = (
        runner.construir_resultado_ok(
            1,
            orden,
            Path("DATA.xlsx"),
            fila_excel=27,
        )
    )

    assert (
        resultado["FILA_EXCEL"]
        == 27
    )


def test_resultado_error_usa_fila_cacheada_sin_abrir_excel(
    monkeypatch,
):
    def prohibido(*args, **kwargs):
        raise AssertionError(
            "No debe abrir Excel para buscar la fila."
        )

    monkeypatch.setattr(
        runner,
        "obtener_fila_excel_por_id",
        prohibido,
    )

    monkeypatch.setattr(
        runner,
        "construir_mensaje_usuario",
        lambda error, orden: "Error prueba",
    )

    orden = SimpleNamespace(
        id_orden="1001"
    )

    resultado = (
        runner.construir_resultado_error(
            1,
            orden,
            Path("DATA.xlsx"),
            RuntimeError("fallo"),
            {},
            fila_excel=27,
        )
    )

    assert (
        resultado["FILA_EXCEL"]
        == 27
    )


def test_resultado_mantiene_fallback_si_no_recibe_fila(
    monkeypatch,
):
    llamadas = {
        "total": 0,
    }

    def buscar(*args, **kwargs):
        llamadas["total"] += 1
        return 44

    monkeypatch.setattr(
        runner,
        "obtener_fila_excel_por_id",
        buscar,
    )

    orden = SimpleNamespace(
        id_orden="1001"
    )

    resultado = (
        runner.construir_resultado_ok(
            1,
            orden,
            Path("DATA.xlsx"),
        )
    )

    assert resultado["FILA_EXCEL"] == 44
    assert llamadas["total"] == 1


def test_ejecutar_rpa_usa_cache_de_filas():
    codigo = inspect.getsource(
        runner.ejecutar_rpa
    )

    assert (
        "filas_excel_por_id"
        in codigo
    )

    assert (
        "fila_excel=fila_excel_orden"
        in codigo
    )

    # La búsqueda histórica queda solo como fallback en
    # _resolver_fila_excel_resultado, no en el loop principal.
    assert (
        "obtener_fila_excel_por_id(ruta_excel, orden.id_orden)"
        not in codigo
    )



# ============================================================
# CP13B.1 - baseline medible de I/O Excel
# ============================================================


def _crear_excel_cp13_io(
    ruta: Path,
) -> Path:
    """
    Workbook mínimo para medir aperturas/guardados del pipeline
    de preparación sin depender de datos reales.
    """
    wb = Workbook()

    try:
        ws = wb.active
        ws.title = "Ordenes"

        ws.append(
            [
                "ID_ORDEN",
                "ESTADO_RPA",
                "RESUMEN",
                "DATO_NEGOCIO",
            ]
        )

        ws.append(
            [
                "1001",
                0,
                "",
                "A",
            ]
        )

        ws.append(
            [
                "1002",
                0,
                "",
                "B",
            ]
        )

        adjuntos = wb.create_sheet(
            "Adjuntos"
        )

        adjuntos.append(
            [
                "ID_ORDEN",
                "ARCHIVO",
            ]
        )

        adjuntos.append(
            [
                "1001",
                "documento.pdf",
            ]
        )

        wb.save(
            ruta
        )

    finally:
        wb.close()

    return ruta


def test_cp13_baseline_preparacion_actual_mide_io(
    tmp_path,
    monkeypatch,
):
    """
    Congela el costo actual de crear_excel_trabajo_pendientes
    usando la ruta tradicional (solo_lectura=False).

    Baseline esperado:
    - 2 load_workbook desde state_manager:
        1 preparar/validar
        1 obtener pendientes
    - 1 escritura atómica del snapshot
    - 1 pd.read_excel de todas las hojas
    """
    excel = _crear_excel_cp13_io(
        tmp_path / "entrada_original.xlsx"
    )

    temporal = (
        tmp_path
        / "temporales"
    )

    monkeypatch.setattr(
        state_manager,
        "obtener_directorio_temporal",
        lambda: temporal,
    )

    contadores = {
        "load_workbook": 0,
        "guardar_atomico": 0,
        "read_excel": 0,
    }

    load_original = (
        state_manager.load_workbook
    )

    guardar_original = (
        state_manager.guardar_workbook_atomico
    )

    read_excel_original = (
        state_manager.pd.read_excel
    )

    def contar_load(
        *args,
        **kwargs,
    ):
        contadores[
            "load_workbook"
        ] += 1

        return load_original(
            *args,
            **kwargs,
        )

    def contar_guardado(
        *args,
        **kwargs,
    ):
        contadores[
            "guardar_atomico"
        ] += 1

        return guardar_original(
            *args,
            **kwargs,
        )

    def contar_read_excel(
        *args,
        **kwargs,
    ):
        contadores[
            "read_excel"
        ] += 1

        return read_excel_original(
            *args,
            **kwargs,
        )

    monkeypatch.setattr(
        state_manager,
        "load_workbook",
        contar_load,
    )

    monkeypatch.setattr(
        state_manager,
        "guardar_workbook_atomico",
        contar_guardado,
    )

    monkeypatch.setattr(
        state_manager.pd,
        "read_excel",
        contar_read_excel,
    )

    ruta_trabajo, pendientes = (
        state_manager.crear_excel_trabajo_pendientes(
            excel
        )
    )

    assert ruta_trabajo.exists()

    assert [
        pendiente.id_orden
        for pendiente in pendientes
    ] == [
        "1001",
        "1002",
    ]

    assert contadores == {
        "load_workbook": 2,
        "guardar_atomico": 1,
        "read_excel": 1,
    }


def test_cp13_baseline_ruta_readonly_reduce_io(
    tmp_path,
    monkeypatch,
):
    """
    Mide la alternativa que ya existe en state_manager.

    Todavía NO cambia runner.py.

    Esta prueba demuestra que inspeccionar pendientes en
    solo lectura evita una apertura y el guardado atómico.
    """
    excel = _crear_excel_cp13_io(
        tmp_path / "entrada_original.xlsx"
    )

    temporal = (
        tmp_path
        / "temporales"
    )

    monkeypatch.setattr(
        state_manager,
        "obtener_directorio_temporal",
        lambda: temporal,
    )

    contadores = {
        "load_workbook": 0,
        "guardar_atomico": 0,
        "read_excel": 0,
    }

    load_original = (
        state_manager.load_workbook
    )

    guardar_original = (
        state_manager.guardar_workbook_atomico
    )

    read_excel_original = (
        state_manager.pd.read_excel
    )

    def contar_load(
        *args,
        **kwargs,
    ):
        contadores[
            "load_workbook"
        ] += 1

        return load_original(
            *args,
            **kwargs,
        )

    def contar_guardado(
        *args,
        **kwargs,
    ):
        contadores[
            "guardar_atomico"
        ] += 1

        return guardar_original(
            *args,
            **kwargs,
        )

    def contar_read_excel(
        *args,
        **kwargs,
    ):
        contadores[
            "read_excel"
        ] += 1

        return read_excel_original(
            *args,
            **kwargs,
        )

    monkeypatch.setattr(
        state_manager,
        "load_workbook",
        contar_load,
    )

    monkeypatch.setattr(
        state_manager,
        "guardar_workbook_atomico",
        contar_guardado,
    )

    monkeypatch.setattr(
        state_manager.pd,
        "read_excel",
        contar_read_excel,
    )

    ruta_trabajo, pendientes = (
        state_manager.crear_excel_trabajo_pendientes(
            excel,
            solo_lectura=True,
        )
    )

    assert ruta_trabajo.exists()

    assert [
        pendiente.id_orden
        for pendiente in pendientes
    ] == [
        "1001",
        "1002",
    ]

    assert contadores == {
        "load_workbook": 1,
        "guardar_atomico": 0,
        "read_excel": 1,
    }


def test_cp13_baseline_versiones_abren_snapshot_una_vez(
    tmp_path,
    monkeypatch,
):
    """
    Capturar todas las versiones de las órdenes usa una única
    apertura read-only del snapshot, independientemente de la
    cantidad de ID_ORDEN solicitados.
    """
    excel = _crear_excel_cp13_io(
        tmp_path / "entrada_original.xlsx"
    )

    llamadas = {
        "load_workbook": 0,
    }

    load_original = (
        concurrency_guard.load_workbook
    )

    def contar_load(
        *args,
        **kwargs,
    ):
        llamadas[
            "load_workbook"
        ] += 1

        return load_original(
            *args,
            **kwargs,
        )

    monkeypatch.setattr(
        concurrency_guard,
        "load_workbook",
        contar_load,
    )

    versiones = (
        concurrency_guard.capturar_versiones_ordenes(
            excel,
            [
                "1001",
                "1002",
            ],
        )
    )

    assert set(
        versiones
    ) == {
        "1001",
        "1002",
    }

    assert llamadas[
        "load_workbook"
    ] == 1


def test_cp13_runner_usa_preparacion_readonly():
    """
    CP13B.2: el snapshot de entrada ya no debe reescribirse
    únicamente para identificar órdenes pendientes.
    """
    codigo = inspect.getsource(
        runner.ejecutar_rpa
    )

    assert (
        "crear_excel_trabajo_pendientes("
        in codigo
    )

    assert (
        "solo_lectura=True"
        in codigo
    )

    assert (
        "crear_excel_trabajo_pendientes(entrada_original)"
        not in codigo
    )



def test_cp13_readonly_no_modifica_snapshot_sin_estado(
    tmp_path,
    monkeypatch,
):
    """
    Un snapshot que todavía no contiene ESTADO_RPA debe poder
    inspeccionarse como pendiente sin modificar físicamente
    el archivo.
    """
    excel = (
        tmp_path
        / "entrada_sin_estado.xlsx"
    )

    wb = Workbook()

    try:
        ws = wb.active
        ws.title = "Ordenes"

        ws.append(
            [
                "ID_ORDEN",
                "RESUMEN",
                "DATO_NEGOCIO",
            ]
        )

        ws.append(
            [
                "1001",
                "",
                "A",
            ]
        )

        ws.append(
            [
                "1002",
                "",
                "B",
            ]
        )

        adjuntos = wb.create_sheet(
            "Adjuntos"
        )

        adjuntos.append(
            [
                "ID_ORDEN",
                "ARCHIVO",
            ]
        )

        adjuntos.append(
            [
                "1001",
                "documento.pdf",
            ]
        )

        wb.save(
            excel
        )

    finally:
        wb.close()

    temporal = (
        tmp_path
        / "temporales"
    )

    monkeypatch.setattr(
        state_manager,
        "obtener_directorio_temporal",
        lambda: temporal,
    )

    contenido_antes = (
        excel.read_bytes()
    )

    ruta_trabajo, pendientes = (
        state_manager.crear_excel_trabajo_pendientes(
            excel,
            solo_lectura=True,
        )
    )

    contenido_despues = (
        excel.read_bytes()
    )

    assert ruta_trabajo.exists()

    assert [
        pendiente.id_orden
        for pendiente in pendientes
    ] == [
        "1001",
        "1002",
    ]

    # Garantía fuerte:
    # el XLSX original queda byte-a-byte idéntico.
    assert (
        contenido_despues
        == contenido_antes
    )


def test_cp13_readonly_respeta_estados_existentes(
    tmp_path,
    monkeypatch,
):
    """
    La optimización no debe convertir en pendiente una orden
    cuyo ESTADO_RPA ya indica que fue procesada.
    """
    excel = (
        tmp_path
        / "entrada_estados.xlsx"
    )

    wb = Workbook()

    try:
        ws = wb.active
        ws.title = "Ordenes"

        ws.append(
            [
                "ID_ORDEN",
                "ESTADO_RPA",
                "RESUMEN",
                "DATO_NEGOCIO",
            ]
        )

        ws.append(
            [
                "1001",
                0,
                "",
                "PENDIENTE",
            ]
        )

        ws.append(
            [
                "1002",
                1,
                "OK",
                "PROCESADA",
            ]
        )

        wb.save(
            excel
        )

    finally:
        wb.close()

    temporal = (
        tmp_path
        / "temporales"
    )

    monkeypatch.setattr(
        state_manager,
        "obtener_directorio_temporal",
        lambda: temporal,
    )

    contenido_antes = (
        excel.read_bytes()
    )

    _, pendientes = (
        state_manager.crear_excel_trabajo_pendientes(
            excel,
            solo_lectura=True,
        )
    )

    assert [
        pendiente.id_orden
        for pendiente in pendientes
    ] == [
        "1001",
    ]

    assert (
        excel.read_bytes()
        == contenido_antes
    )
