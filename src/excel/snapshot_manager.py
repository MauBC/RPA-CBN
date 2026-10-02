from __future__ import annotations

from dataclasses import replace
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any
import hashlib
import json
import os
import shutil
import time
import uuid
import zipfile

from openpyxl import load_workbook

from src.excel.validators import OrdenCotizacion


class SnapshotError(RuntimeError):
    """Error al crear una copia consistente de un input del RPA."""


def sha256_archivo(ruta: str | Path) -> str:
    ruta = Path(ruta)

    digest = hashlib.sha256()

    with ruta.open("rb") as archivo:
        while True:
            bloque = archivo.read(1024 * 1024)

            if not bloque:
                break

            digest.update(bloque)

    return digest.hexdigest()


def crear_snapshot_estable(
    origen: str | Path,
    destino: str | Path,
    *,
    intentos: int = 3,
    espera_segundos: float = 0.15,
) -> dict[str, Any]:
    """
    Copia un archivo asegurando que el origen no cambió durante la copia.

    Para cada intento:
    1. SHA256 origen antes.
    2. Copia a temporal.
    3. SHA256 copia.
    4. SHA256 origen después.
    5. Solo acepta si los tres hashes coinciden.
    """
    origen = Path(origen).resolve()
    destino = Path(destino).resolve()

    if not origen.exists():
        raise SnapshotError(
            f"No existe el archivo a capturar: {origen}"
        )

    if not origen.is_file():
        raise SnapshotError(
            f"La ruta no es un archivo: {origen}"
        )

    if origen == destino:
        raise SnapshotError(
            "El snapshot no puede sobrescribir el archivo de origen."
        )

    destino.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    ultimo_error: Exception | None = None

    for intento in range(1, intentos + 1):
        temporal = destino.with_name(
            f".{destino.name}."
            f"{os.getpid()}."
            f"{uuid.uuid4().hex}.tmp"
        )

        try:
            hash_antes = sha256_archivo(origen)

            shutil.copy2(
                origen,
                temporal,
            )

            hash_copia = sha256_archivo(temporal)
            hash_despues = sha256_archivo(origen)

            if (
                hash_antes
                == hash_copia
                == hash_despues
            ):
                os.replace(
                    temporal,
                    destino,
                )

                stat = origen.stat()

                return {
                    "origen": str(origen),
                    "snapshot": str(destino),
                    "sha256": hash_copia,
                    "bytes": destino.stat().st_size,
                    "origen_mtime_ns": stat.st_mtime_ns,
                    "capturado_en": datetime.now().isoformat(
                        timespec="seconds"
                    ),
                    "intento": intento,
                }

            ultimo_error = SnapshotError(
                f"El archivo cambió durante la copia "
                f"(intento {intento}/{intentos})."
            )

        except (
            PermissionError,
            OSError,
        ) as exc:
            ultimo_error = exc

        finally:
            try:
                temporal.unlink()
            except FileNotFoundError:
                pass
            except Exception:
                pass

        if intento < intentos:
            time.sleep(espera_segundos)

    detalle = (
        f" Detalle: {ultimo_error}"
        if ultimo_error is not None
        else ""
    )

    raise SnapshotError(
        f"No se pudo crear un snapshot consistente de "
        f"'{origen}' después de {intentos} intentos."
        f"{detalle}"
    )



def _quantity_or_percent_a_numero_excel(
    valor: Any,
    *,
    contexto: str,
) -> int | float:
    """
    Convierte quantityOrPercent al mismo valor numérico que
    acepta actualmente el lector del RPA.

    Es intencionalmente compatible con la semántica existente:
    - espacios se ignoran;
    - coma se interpreta como separador de miles;
    - se redondea a 4 decimales con ROUND_HALF_UP;
    - solo admite valores mayores a cero.
    """
    if valor is None:
        raise SnapshotError(
            f"{contexto}: quantityOrPercent vacío."
        )

    texto = str(
        valor
    ).strip()

    if not texto:
        raise SnapshotError(
            f"{contexto}: quantityOrPercent vacío."
        )

    texto = texto.replace(
        ",",
        "",
    )

    try:
        numero = Decimal(
            texto
        )
    except InvalidOperation as exc:
        raise SnapshotError(
            f"{contexto}: quantityOrPercent "
            f"no numérico: {valor!r}."
        ) from exc

    if (
        not numero.is_finite()
        or numero <= 0
    ):
        raise SnapshotError(
            f"{contexto}: quantityOrPercent "
            f"debe ser mayor a cero. "
            f"Valor: {valor!r}."
        )

    numero = numero.quantize(
        Decimal("0.0001"),
        rounding=ROUND_HALF_UP,
    )

    entero = numero.to_integral_value()

    if numero == entero:
        return int(
            entero
        )

    return float(
        numero
    )


def normalizar_quantity_or_percent_snapshot(
    ruta_template: str | Path,
) -> dict[str, Any]:
    """
    Normaliza exclusivamente el TEMPLATE congelado que será
    entregado al portal.

    El archivo original del usuario nunca se modifica.

    quantityOrPercent queda guardado físicamente como tipo
    numérico de Excel, incluso cuando el origen contenía:
        "300"
        " 300 "
        "300.00"
        formato de celda "@"

    Si existen fórmulas en quantityOrPercent, no reescribe el
    workbook. Se conserva el comportamiento anterior para no
    destruir valores calculados/caché de Excel.
    """
    ruta_template = Path(
        ruta_template
    ).resolve()

    keep_vba = (
        ruta_template.suffix.lower()
        == ".xlsm"
    )

    try:
        wb = load_workbook(
            ruta_template,
            read_only=False,
            data_only=False,
            keep_vba=keep_vba,
        )
    except Exception as exc:
        raise SnapshotError(
            "No se pudo abrir el TEMPLATE congelado "
            f"para normalización: {ruta_template}"
        ) from exc

    temporal: Path | None = None

    try:
        if len(
            wb.sheetnames
        ) != 1:
            raise SnapshotError(
                "TEMPLATE congelado debe tener "
                "exactamente una hoja. "
                f"Archivo: {ruta_template}"
            )

        ws = wb[
            wb.sheetnames[0]
        ]

        header_a = str(
            ws["A1"].value
            or ""
        ).strip()

        header_b = str(
            ws["B1"].value
            or ""
        ).strip()

        if (
            header_a != "code"
            or header_b
            != "quantityOrPercent"
        ):
            raise SnapshotError(
                "TEMPLATE congelado tiene cabeceras "
                "inválidas. "
                f"Archivo: {ruta_template}"
            )

        formulas = []

        for fila in range(
            3,
            ws.max_row + 1,
        ):
            celda = ws.cell(
                row=fila,
                column=2,
            )

            if celda.data_type == "f":
                formulas.append(
                    fila
                )

        if formulas:
            return {
                "modificado": False,
                "normalizadas": 0,
                "formulas_omitidas": len(
                    formulas
                ),
            }

        normalizadas = 0

        for fila in range(
            3,
            ws.max_row + 1,
        ):
            celda_codigo = ws.cell(
                row=fila,
                column=1,
            )

            celda_valor = ws.cell(
                row=fila,
                column=2,
            )

            codigo = celda_codigo.value
            valor = celda_valor.value

            codigo_vacio = (
                codigo is None
                or not str(
                    codigo
                ).strip()
            )

            valor_vacio = (
                valor is None
                or (
                    isinstance(
                        valor,
                        str,
                    )
                    and not valor.strip()
                )
            )

            if (
                codigo_vacio
                and valor_vacio
            ):
                continue

            if valor_vacio:
                raise SnapshotError(
                    f"TEMPLATE {ruta_template}, "
                    f"fila {fila}: "
                    "quantityOrPercent vacío."
                )

            numero_excel = (
                _quantity_or_percent_a_numero_excel(
                    valor,
                    contexto=(
                        f"TEMPLATE {ruta_template}, "
                        f"fila {fila}"
                    ),
                )
            )

            debe_normalizar = (
                isinstance(
                    valor,
                    str,
                )
                or celda_valor.data_type
                != "n"
                or celda_valor.number_format
                == "@"
            )

            if debe_normalizar:
                normalizadas += 1

            # Se asigna siempre un tipo Python numérico.
            # OpenPyXL lo serializa como celda Excel tipo "n".
            celda_valor.value = (
                numero_excel
            )

            # Evita conservar formato explícito "Texto".
            if (
                celda_valor.number_format
                == "@"
            ):
                celda_valor.number_format = (
                    "General"
                )

        if normalizadas == 0:
            return {
                "modificado": False,
                "normalizadas": 0,
                "formulas_omitidas": 0,
            }

        temporal = (
            ruta_template.with_name(
                f".{ruta_template.stem}."
                f"{os.getpid()}."
                f"{uuid.uuid4().hex}"
                f"{ruta_template.suffix}"
            )
        )

        wb.save(
            temporal
        )

        wb.close()

        wb = None

        os.replace(
            temporal,
            ruta_template,
        )

        temporal = None

        return {
            "modificado": True,
            "normalizadas": normalizadas,
            "formulas_omitidas": 0,
        }

    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass

        if temporal is not None:
            try:
                temporal.unlink()
            except FileNotFoundError:
                pass
            except Exception:
                pass


def crear_snapshot_template_estable(
    origen: str | Path,
    destino: str | Path,
) -> dict[str, Any]:
    """
    Snapshot estable + normalización del archivo que realmente
    se subirá a CBN.

    La metadata final describe el snapshot ya normalizado.
    El SHA original también se conserva para trazabilidad.
    """
    metadata = crear_snapshot_estable(
        origen,
        destino,
    )

    sha256_origen = metadata[
        "sha256"
    ]

    destino = Path(
        destino
    ).resolve()

    # Históricamente este helper también puede usarse con
    # fixtures/opacos que tienen extensión .xlsx pero no son
    # realmente un workbook ZIP. En ese caso conservamos el
    # comportamiento anterior: snapshot exacto, sin intentar
    # interpretar su contenido.
    #
    # En el flujo productivo los TEMPLATE reales ya fueron
    # validados por reader.py antes de llegar aquí.
    if zipfile.is_zipfile(
        destino
    ):
        normalizacion = (
            normalizar_quantity_or_percent_snapshot(
                destino
            )
        )
    else:
        normalizacion = {
            "modificado": False,
            "normalizadas": 0,
            "formulas_omitidas": 0,
            "omitido_no_excel": True,
            "motivo": "snapshot_no_es_xlsx_zip",
        }

    metadata[
        "sha256_origen"
    ] = sha256_origen

    metadata[
        "sha256"
    ] = sha256_archivo(
        destino
    )

    metadata[
        "bytes"
    ] = destino.stat().st_size

    metadata[
        "normalizacion_quantity_or_percent"
    ] = normalizacion

    return metadata


def _clave_ruta(ruta: Path) -> str:
    return os.path.normcase(
        str(ruta.resolve())
    )


def _nombre_snapshot_template(
    ruta: Path,
) -> str:
    clave = _clave_ruta(ruta)

    prefijo = hashlib.sha256(
        clave.encode("utf-8")
    ).hexdigest()[:10]

    # Conservamos el nombre original.
    # Esto también conserva reglas como "_porcent_".
    return f"{prefijo}_{ruta.name}"


def congelar_templates_ordenes(
    ordenes: list[OrdenCotizacion],
    directorio_templates: str | Path,
) -> tuple[list[OrdenCotizacion], list[dict[str, Any]]]:
    """
    Copia los TEMPLATE utilizados por las órdenes y devuelve nuevas
    dataclasses apuntando exclusivamente a los snapshots.

    Si varias posiciones usan el mismo template, solo se copia una vez.
    """
    directorio_templates = Path(
        directorio_templates
    )

    directorio_templates.mkdir(
        parents=True,
        exist_ok=True,
    )

    snapshots_por_origen: dict[str, Path] = {}
    metadatos: list[dict[str, Any]] = []

    def obtener_snapshot(
        ruta: Path | None,
    ) -> Path | None:
        if ruta is None:
            return None

        ruta = Path(ruta).resolve()
        clave = _clave_ruta(ruta)

        existente = snapshots_por_origen.get(
            clave
        )

        if existente is not None:
            return existente

        destino = (
            directorio_templates
            / _nombre_snapshot_template(ruta)
        )

        metadata = crear_snapshot_template_estable(
            ruta,
            destino,
        )

        metadata["tipo"] = "template"

        snapshots_por_origen[clave] = destino
        metadatos.append(metadata)

        return destino

    ordenes_congeladas: list[
        OrdenCotizacion
    ] = []

    for orden in ordenes:
        posiciones_nuevas = []

        for posicion in orden.posiciones_detalle:
            posiciones_nuevas.append(
                replace(
                    posicion,
                    template=obtener_snapshot(
                        posicion.template
                    ),
                )
            )

        ordenes_congeladas.append(
            replace(
                orden,
                template=obtener_snapshot(
                    orden.template
                ),
                posiciones_detalle=posiciones_nuevas,
            )
        )

    return ordenes_congeladas, metadatos


def guardar_manifest_snapshots(
    ruta_manifest: str | Path,
    entrada: dict[str, Any],
    templates: list[dict[str, Any]],
) -> Path:
    ruta_manifest = Path(ruta_manifest)

    ruta_manifest.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    datos = {
        "version": 1,
        "generado_en": datetime.now().isoformat(
            timespec="seconds"
        ),
        "entrada": entrada,
        "templates": templates,
    }

    temporal = ruta_manifest.with_name(
        f".{ruta_manifest.name}."
        f"{os.getpid()}."
        f"{uuid.uuid4().hex}.tmp"
    )

    try:
        temporal.write_text(
            json.dumps(
                datos,
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )

        os.replace(
            temporal,
            ruta_manifest,
        )

    finally:
        try:
            temporal.unlink()
        except FileNotFoundError:
            pass
        except Exception:
            pass

    return ruta_manifest
