from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable
import hashlib
import json

from openpyxl import load_workbook


HOJA_ORDENES = "Ordenes"
HOJA_ADJUNTOS = "Adjuntos"

COL_ID_ORDEN = "ID_ORDEN"
COL_ESTADO_RPA = "ESTADO_RPA"
COL_RESUMEN = "RESUMEN"

COLUMNAS_CONTROL_RPA = {
    COL_ESTADO_RPA,
    COL_RESUMEN,
    "_FILA_EXCEL_ORIGINAL",
}


class ExcelConflictoEdicionError(ValueError):
    """
    El DATA.xlsx ya no representa la misma orden
    que fue congelada para la ejecucion actual.
    """


def _normalizar_header(
    valor: Any,
) -> str:
    return str(
        valor or ""
    ).strip().upper()


def _normalizar_id(
    valor: Any,
) -> str:
    if valor is None:
        return ""

    texto = str(
        valor
    ).strip()

    if texto.endswith(".0"):
        texto = texto[:-2]

    return texto


def _normalizar_estado(
    valor: Any,
) -> int:
    if valor is None:
        return 0

    texto = str(
        valor
    ).strip()

    if not texto:
        return 0

    if texto.endswith(".0"):
        texto = texto[:-2]

    try:
        return int(
            texto
        )
    except Exception:
        return 0


def _normalizar_resumen(
    valor: Any,
) -> str | None:
    if valor is None:
        return None

    texto = str(
        valor
    ).strip()

    return (
        texto
        if texto
        else None
    )


def _serializar_valor(
    valor: Any,
) -> Any:
    if valor is None:
        return None

    if isinstance(
        valor,
        datetime,
    ):
        return valor.isoformat()

    if isinstance(
        valor,
        date,
    ):
        return valor.isoformat()

    if isinstance(
        valor,
        Decimal,
    ):
        return format(
            valor,
            "f",
        )

    if isinstance(
        valor,
        (
            str,
            int,
            float,
            bool,
        ),
    ):
        return valor

    return str(
        valor
    )


def _headers_ws(
    ws,
) -> list[tuple[int, str]]:
    resultado = []

    for cell in ws[1]:
        header = _normalizar_header(
            cell.value
        )

        if header:
            resultado.append(
                (
                    cell.column,
                    header,
                )
            )

    return resultado


def _capturar_filas_hoja(
    ws,
    id_orden: str,
    *,
    capturar_control: bool,
) -> tuple[
    list[list[list[Any]]],
    list[int],
    list[str | None],
]:
    headers = _headers_ws(
        ws
    )

    col_id = None
    col_estado = None
    col_resumen = None

    for columna, header in headers:
        if header == COL_ID_ORDEN:
            col_id = columna
        elif header == COL_ESTADO_RPA:
            col_estado = columna
        elif header == COL_RESUMEN:
            col_resumen = columna

    if col_id is None:
        return [], [], []

    filas_contenido = []
    estados = []
    resumenes = []

    for fila in range(
        2,
        ws.max_row + 1,
    ):
        actual = _normalizar_id(
            ws.cell(
                fila,
                col_id,
            ).value
        )

        if actual != id_orden:
            continue

        contenido_fila = []

        for columna, header in headers:
            if (
                header
                in COLUMNAS_CONTROL_RPA
            ):
                continue

            valor = _serializar_valor(
                ws.cell(
                    fila,
                    columna,
                ).value
            )

            # Una columna nueva completamente vacia
            # no debe crear un falso conflicto.
            if (
                valor is None
                or valor == ""
            ):
                continue

            contenido_fila.append(
                [
                    header,
                    valor,
                ]
            )

        filas_contenido.append(
            contenido_fila
        )

        if capturar_control:
            estado = (
                _normalizar_estado(
                    ws.cell(
                        fila,
                        col_estado,
                    ).value
                )
                if col_estado is not None
                else 0
            )

            resumen = (
                _normalizar_resumen(
                    ws.cell(
                        fila,
                        col_resumen,
                    ).value
                )
                if col_resumen is not None
                else None
            )

            estados.append(
                estado
            )

            resumenes.append(
                resumen
            )

    return (
        filas_contenido,
        estados,
        resumenes,
    )


def capturar_version_orden_workbook(
    workbook,
    id_orden: Any,
) -> dict[str, Any]:
    id_normalizado = _normalizar_id(
        id_orden
    )

    if not id_normalizado:
        raise ValueError(
            "id_orden no puede estar vacio."
        )

    if (
        HOJA_ORDENES
        not in workbook.sheetnames
    ):
        raise ValueError(
            "No existe la hoja Ordenes."
        )

    (
        filas_ordenes,
        estados,
        resumenes,
    ) = _capturar_filas_hoja(
        workbook[
            HOJA_ORDENES
        ],
        id_normalizado,
        capturar_control=True,
    )

    if not filas_ordenes:
        raise ValueError(
            f"No existe ID_ORDEN={id_normalizado}."
        )

    filas_adjuntos = []

    if (
        HOJA_ADJUNTOS
        in workbook.sheetnames
    ):
        (
            filas_adjuntos,
            _,
            _,
        ) = _capturar_filas_hoja(
            workbook[
                HOJA_ADJUNTOS
            ],
            id_normalizado,
            capturar_control=False,
        )

    payload = {
        "ordenes": filas_ordenes,
        "adjuntos": filas_adjuntos,
    }

    canonical = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(
            ",",
            ":",
        ),
        default=str,
    )

    huella = hashlib.sha256(
        canonical.encode(
            "utf-8"
        )
    ).hexdigest()

    return {
        "id_orden": id_normalizado,
        "huella_contenido": huella,
        "estados_rpa": estados,
        "resumenes": resumenes,
        "filas_ordenes": len(
            filas_ordenes
        ),
        "filas_adjuntos": len(
            filas_adjuntos
        ),
    }


def capturar_versiones_ordenes(
    ruta_excel: str | Path,
    ids_orden: Iterable[Any],
) -> dict[str, dict[str, Any]]:
    ruta_excel = Path(
        ruta_excel
    )

    ids = []

    vistos = set()

    for valor in ids_orden:
        normalizado = _normalizar_id(
            valor
        )

        if (
            normalizado
            and normalizado not in vistos
        ):
            vistos.add(
                normalizado
            )
            ids.append(
                normalizado
            )

    if not ids:
        return {}

    wb = load_workbook(
        ruta_excel,
        read_only=True,
        data_only=False,
    )

    try:
        return {
            id_orden: (
                capturar_version_orden_workbook(
                    wb,
                    id_orden,
                )
            )
            for id_orden in ids
        }

    finally:
        wb.close()


def validar_version_orden_workbook(
    workbook,
    id_orden: Any,
    version_esperada: dict[str, Any],
    *,
    estado_objetivo: int,
    resumen_objetivo: str | None,
) -> None:
    actual = (
        capturar_version_orden_workbook(
            workbook,
            id_orden,
        )
    )

    esperada = dict(
        version_esperada
    )

    if (
        str(
            actual[
                "huella_contenido"
            ]
        )
        != str(
            esperada.get(
                "huella_contenido",
                "",
            )
        )
    ):
        raise ExcelConflictoEdicionError(
            "La orden fue modificada en DATA.xlsx "
            "despues de iniciar la ejecucion. "
            f"ID_ORDEN={actual['id_orden']}."
        )

    estados_actuales = list(
        actual.get(
            "estados_rpa",
            [],
        )
    )

    estados_esperados = list(
        esperada.get(
            "estados_rpa",
            [],
        )
    )

    if (
        estados_actuales
        != estados_esperados
        and not (
            estados_actuales
            and all(
                estado
                == estado_objetivo
                for estado
                in estados_actuales
            )
        )
    ):
        raise ExcelConflictoEdicionError(
            "ESTADO_RPA cambio durante la ejecucion. "
            f"ID_ORDEN={actual['id_orden']}."
        )

    if resumen_objetivo is None:
        return

    resumen_objetivo = str(
        resumen_objetivo
    ).strip()

    resumenes_actuales = list(
        actual.get(
            "resumenes",
            [],
        )
    )

    resumenes_esperados = list(
        esperada.get(
            "resumenes",
            [],
        )
    )

    if (
        resumenes_actuales
        != resumenes_esperados
        and not (
            resumenes_actuales
            and all(
                (
                    valor
                    or ""
                )
                == resumen_objetivo
                for valor
                in resumenes_actuales
            )
        )
    ):
        raise ExcelConflictoEdicionError(
            "RESUMEN cambio durante la ejecucion. "
            f"ID_ORDEN={actual['id_orden']}."
        )
