from __future__ import annotations

from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from typing import Any
import os
import time

import pandas as pd
from openpyxl import load_workbook

from src.utils.app_paths import obtener_directorio_temporal
from src.excel.excel_transaction import (
    ExcelArchivoCambioConcurrenteError,
    exigir_excel_disponible,
    guardar_workbook_atomico,
    sha256_archivo_excel,
)
from src.excel.concurrency_guard import (
    validar_version_orden_workbook,
)


HOJA_ORDENES = "Ordenes"
COL_ID_ORDEN = "ID_ORDEN"
COL_ESTADO_RPA = "ESTADO_RPA"
COL_RESUMEN = "RESUMEN"


@dataclass(frozen=True)
class PendienteRPA:
    id_orden: str
    fila_excel: int
    estado_rpa: int


def _normalizar_header(valor: Any) -> str:
    return str(valor or "").strip().upper()


def _normalizar_id(valor: Any) -> str:
    if valor is None:
        return ""

    try:
        if pd.isna(valor):
            return ""
    except Exception:
        pass

    texto = str(valor).strip()

    if texto.endswith(".0"):
        texto = texto[:-2]

    return texto


def _estado_a_int(valor: Any) -> int | None:
    if valor is None:
        return None

    try:
        if pd.isna(valor):
            return None
    except Exception:
        pass

    texto = str(valor).strip()

    if texto.endswith(".0"):
        texto = texto[:-2]

    try:
        return int(texto)
    except Exception:
        return None


def _headers_ws(ws) -> dict[str, int]:
    headers: dict[str, int] = {}

    for cell in ws[1]:
        nombre = _normalizar_header(cell.value)

        if nombre:
            headers[nombre] = cell.column

    return headers


def _env_int(nombre: str) -> int | None:
    valor = os.getenv(nombre, "").strip()

    if not valor:
        return None

    try:
        return int(valor)
    except Exception:
        return None


def _worker_config() -> tuple[int | None, int | None]:
    workers = _env_int("RPA_WORKERS")
    worker_id = _env_int("RPA_WORKER_ID")

    if not workers or not worker_id or workers <= 1:
        return None, None

    if worker_id < 1 or worker_id > workers:
        raise ValueError(
            f"Config worker invalida: RPA_WORKER_ID={worker_id}, RPA_WORKERS={workers}."
        )

    return worker_id, workers


def _id_o_fila_para_paridad(pendiente: PendienteRPA) -> int:
    try:
        return int(float(str(pendiente.id_orden).strip()))
    except Exception:
        return int(pendiente.fila_excel)


def _filtrar_pendientes_por_worker(pendientes: list[PendienteRPA]) -> list[PendienteRPA]:
    worker_id, workers = _worker_config()

    if not worker_id or not workers:
        return pendientes

    filtradas: list[PendienteRPA] = []

    for indice, pendiente in enumerate(pendientes, start=1):
        if workers == 2:
            numero = _id_o_fila_para_paridad(pendiente)

            # Worker 1 = impares, Worker 2 = pares.
            if worker_id == 1 and numero % 2 == 1:
                filtradas.append(pendiente)
            elif worker_id == 2 and numero % 2 == 0:
                filtradas.append(pendiente)

        else:
            if ((indice - 1) % workers) + 1 == worker_id:
                filtradas.append(pendiente)

    print(
        f"Worker {worker_id}/{workers}: "
        f"{len(filtradas)} orden(es) asignadas de {len(pendientes)} pendiente(s)."
    )

    if workers == 2:
        if worker_id == 1:
            print("Worker 1 procesa ID_ORDEN impares.")
        elif worker_id == 2:
            print("Worker 2 procesa ID_ORDEN pares.")

    return filtradas


def _sufijo_worker() -> str:
    worker_id, workers = _worker_config()

    if worker_id and workers:
        return f"_worker_{worker_id}"

    return ""


@contextmanager
def _excel_lock(ruta_excel: str | Path, timeout_segundos: int = 180):
    ruta_excel = Path(ruta_excel)
    ruta_lock = Path(str(ruta_excel) + ".lock")
    inicio = time.time()
    fd = None

    while True:
        try:
            fd = os.open(str(ruta_lock), os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.write(fd, f"pid={os.getpid()} time={time.time()}".encode("utf-8"))
            break

        except FileExistsError:
            try:
                if ruta_lock.exists() and time.time() - ruta_lock.stat().st_mtime > 600:
                    ruta_lock.unlink()
                    continue
            except Exception:
                pass

            if time.time() - inicio > timeout_segundos:
                raise TimeoutError(f"No se pudo obtener lock de Excel: {ruta_lock}")

            time.sleep(0.5)

    try:
        yield

    finally:
        if fd is not None:
            try:
                os.close(fd)
            except Exception:
                pass

        try:
            ruta_lock.unlink()
        except FileNotFoundError:
            pass
        except Exception:
            pass


def _cargar_workbook_estable_para_actualizacion(
    ruta_excel: str | Path,
    *,
    intentos: int = 3,
) -> tuple[Any, str]:
    """
    Carga el workbook ?nicamente cuando DATA.xlsx
    permanece idéntico durante toda la lectura.

    Devuelve:
        (workbook, sha256_version_cargada)
    """
    ruta_excel = Path(
        ruta_excel
    )

    ultimo_sha_antes = ""
    ultimo_sha_despues = ""

    for intento in range(
        1,
        intentos + 1,
    ):
        exigir_excel_disponible(
            ruta_excel
        )

        sha_antes = (
            sha256_archivo_excel(
                ruta_excel
            )
        )

        wb = load_workbook(
            ruta_excel
        )

        try:
            sha_despues = (
                sha256_archivo_excel(
                    ruta_excel
                )
            )

            if (
                sha_antes
                == sha_despues
            ):
                return (
                    wb,
                    sha_despues,
                )

        except Exception:
            wb.close()
            raise

        wb.close()

        ultimo_sha_antes = sha_antes
        ultimo_sha_despues = sha_despues

        if intento < intentos:
            time.sleep(
                0.05
            )

    raise ExcelArchivoCambioConcurrenteError(
        ruta_excel,
        ultimo_sha_antes,
        ultimo_sha_despues,
    )


def _asegurar_columna_estado(ws) -> int:
    headers = _headers_ws(ws)

    if COL_ESTADO_RPA in headers:
        return headers[COL_ESTADO_RPA]

    nueva_col = ws.max_column + 1
    ws.cell(row=1, column=nueva_col).value = COL_ESTADO_RPA

    for fila in range(2, ws.max_row + 1):
        ws.cell(row=fila, column=nueva_col).value = 0

    return nueva_col




def _asegurar_columna_resumen(ws) -> int:
    headers = _headers_ws(ws)

    if COL_RESUMEN in headers:
        return headers[COL_RESUMEN]

    nueva_col = ws.max_column + 1
    ws.cell(
        row=1,
        column=nueva_col,
    ).value = COL_RESUMEN

    return nueva_col


def preparar_columna_estado_y_validar_ids(ruta_excel: str | Path) -> None:
    """
    ID_ORDEN repetidos están permitidos porque representan posiciones.

    Todas las filas del mismo ID_ORDEN deben tener el mismo ESTADO_RPA.
    """
    ruta_excel = Path(ruta_excel)

    with _excel_lock(ruta_excel):
        wb, sha256_cargado = (
            _cargar_workbook_estable_para_actualizacion(
                ruta_excel
            )
        )

        try:
            if HOJA_ORDENES not in wb.sheetnames:
                raise ValueError(
                    f"No existe la hoja '{HOJA_ORDENES}' en {ruta_excel}."
                )

            ws = wb[HOJA_ORDENES]
            headers = _headers_ws(ws)

            if COL_ID_ORDEN not in headers:
                raise ValueError(
                    f"No existe la columna '{COL_ID_ORDEN}' "
                    f"en la hoja '{HOJA_ORDENES}'."
                )

            col_id = headers[COL_ID_ORDEN]
            col_estado = _asegurar_columna_estado(ws)

            estados_por_id: dict[str, list[tuple[int, int | None]]] = {}

            for fila in range(2, ws.max_row + 1):
                id_orden = _normalizar_id(
                    ws.cell(row=fila, column=col_id).value
                )

                if not id_orden:
                    continue

                estado = _estado_a_int(
                    ws.cell(row=fila, column=col_estado).value
                )

                estados_por_id.setdefault(
                    id_orden,
                    [],
                ).append((fila, estado))

            inconsistencias: list[str] = []

            for id_orden, datos in estados_por_id.items():
                estados = {
                    estado
                    for _, estado in datos
                }

                if len(estados) > 1:
                    detalle = ", ".join(
                        f"fila {fila}={estado}"
                        for fila, estado in datos
                    )

                    inconsistencias.append(
                        f"ID_ORDEN={id_orden}: {detalle}"
                    )

            if inconsistencias:
                raise ValueError(
                    "Las filas de un mismo ID_ORDEN deben tener "
                    "el mismo ESTADO_RPA:\n"
                    + "\n".join(inconsistencias)
                )

            guardar_workbook_atomico(
                wb,
                ruta_excel,
                sha256_esperado=sha256_cargado,
            )

        finally:
            wb.close()


def inspeccionar_pendientes(
    ruta_excel: str | Path,
) -> list[PendienteRPA]:
    """
    Inspecciona las órdenes pendientes SIN modificar el Excel original.

    Reglas:
    - Si ESTADO_RPA no existe, se considera estado 0 en memoria.
    - Si ESTADO_RPA existe, se respetan sus valores actuales.
    - ID_ORDEN repetidos deben mantener el mismo estado.
    - Nunca agrega columnas.
    - Nunca ejecuta wb.save().
    """
    ruta_excel = Path(ruta_excel)

    wb = load_workbook(
        ruta_excel,
        read_only=True,
        data_only=True,
    )

    try:
        if HOJA_ORDENES not in wb.sheetnames:
            raise ValueError(
                f"No existe la hoja '{HOJA_ORDENES}' "
                f"en {ruta_excel}."
            )

        ws = wb[HOJA_ORDENES]
        headers = _headers_ws(ws)

        if COL_ID_ORDEN not in headers:
            raise ValueError(
                f"No existe la columna '{COL_ID_ORDEN}' "
                f"en la hoja '{HOJA_ORDENES}'."
            )

        col_id = headers[COL_ID_ORDEN]
        col_estado = headers.get(COL_ESTADO_RPA)

        estados_por_id: dict[
            str,
            list[tuple[int, int | None]],
        ] = {}

        pendientes: list[PendienteRPA] = []
        ids_vistos: set[str] = set()

        for fila in range(2, ws.max_row + 1):
            id_orden = _normalizar_id(
                ws.cell(
                    row=fila,
                    column=col_id,
                ).value
            )

            if not id_orden:
                continue

            if col_estado is None:
                # ESTADO_RPA todavía no existe.
                # Durante VALIDACIÓN lo simulamos como pendiente,
                # sin modificar el workbook original.
                estado = 0
            else:
                estado = _estado_a_int(
                    ws.cell(
                        row=fila,
                        column=col_estado,
                    ).value
                )

            estados_por_id.setdefault(
                id_orden,
                [],
            ).append((fila, estado))

            if id_orden not in ids_vistos:
                ids_vistos.add(id_orden)

                if estado == 0:
                    pendientes.append(
                        PendienteRPA(
                            id_orden=id_orden,
                            fila_excel=fila,
                            estado_rpa=0,
                        )
                    )

        inconsistencias: list[str] = []

        for id_orden, datos in estados_por_id.items():
            estados = {
                estado
                for _, estado in datos
            }

            if len(estados) > 1:
                detalle = ", ".join(
                    f"fila {fila}={estado}"
                    for fila, estado in datos
                )

                inconsistencias.append(
                    f"ID_ORDEN={id_orden}: {detalle}"
                )

        if inconsistencias:
            raise ValueError(
                "Las filas de un mismo ID_ORDEN deben tener "
                "el mismo ESTADO_RPA:\n"
                + "\n".join(inconsistencias)
            )

        return _filtrar_pendientes_por_worker(
            pendientes
        )

    finally:
        wb.close()


def obtener_pendientes(ruta_excel: str | Path) -> list[PendienteRPA]:
    ruta_excel = Path(ruta_excel)

    preparar_columna_estado_y_validar_ids(ruta_excel)

    wb = load_workbook(ruta_excel, data_only=True)

    try:
        ws = wb[HOJA_ORDENES]

        headers = _headers_ws(ws)
        col_id = headers[COL_ID_ORDEN]
        col_estado = headers[COL_ESTADO_RPA]

        pendientes: list[PendienteRPA] = []
        ids_vistos: set[str] = set()

        for fila in range(2, ws.max_row + 1):
            id_orden = _normalizar_id(
                ws.cell(row=fila, column=col_id).value
            )

            if not id_orden or id_orden in ids_vistos:
                continue

            ids_vistos.add(id_orden)

            estado = _estado_a_int(
                ws.cell(row=fila, column=col_estado).value
            )

            if estado == 0:
                pendientes.append(
                    PendienteRPA(
                        id_orden=id_orden,
                        fila_excel=fila,
                        estado_rpa=estado,
                    )
                )

        return _filtrar_pendientes_por_worker(pendientes)

    finally:
        wb.close()

def obtener_fila_excel_por_id(
    ruta_excel: str | Path,
    id_orden_buscado: Any,
) -> int | None:
    ruta_excel = Path(ruta_excel)
    id_orden_buscado = _normalizar_id(id_orden_buscado)

    wb = load_workbook(ruta_excel, data_only=True)

    try:
        ws = wb[HOJA_ORDENES]

        headers = _headers_ws(ws)
        col_id = headers[COL_ID_ORDEN]

        for fila in range(2, ws.max_row + 1):
            id_orden = _normalizar_id(
                ws.cell(row=fila, column=col_id).value
            )

            if id_orden == id_orden_buscado:
                return fila

        return None

    finally:
        wb.close()


def actualizar_resultado_orden(
    ruta_excel: str | Path,
    id_orden_buscado: Any,
    *,
    estado_rpa: int,
    resumen: str | None = None,
    version_esperada: dict[str, Any] | None = None,
) -> int:
    """
    Actualiza en una única apertura/guardado los campos controlados
    por el RPA para todas las filas de un ID_ORDEN.

    - ESTADO_RPA siempre se actualiza.
    - RESUMEN solo se actualiza cuando resumen no es None.
    - Todas las posiciones del mismo ID_ORDEN reciben los mismos valores.

    Retorna la cantidad de filas actualizadas.
    """
    if estado_rpa not in (0, 1, 2):
        raise ValueError(
            "estado_rpa solo puede ser 0, 1 o 2."
        )

    ruta_excel = Path(ruta_excel)

    id_orden_buscado = _normalizar_id(
        id_orden_buscado
    )

    if not id_orden_buscado:
        raise ValueError(
            "id_orden_buscado no puede estar vacío."
        )

    texto_resumen = (
        None
        if resumen is None
        else str(resumen).strip()
    )

    with _excel_lock(ruta_excel):
        wb, sha256_cargado = (
            _cargar_workbook_estable_para_actualizacion(
                ruta_excel
            )
        )

        try:
            if HOJA_ORDENES not in wb.sheetnames:
                raise ValueError(
                    f"No existe la hoja '{HOJA_ORDENES}'."
                )

            ws = wb[HOJA_ORDENES]
            headers = _headers_ws(ws)

            if COL_ID_ORDEN not in headers:
                raise ValueError(
                    f"No existe la columna "
                    f"'{COL_ID_ORDEN}'."
                )

            col_id = headers[
                COL_ID_ORDEN
            ]

            if version_esperada is not None:
                validar_version_orden_workbook(
                    wb,
                    id_orden_buscado,
                    version_esperada,
                    estado_objetivo=estado_rpa,
                    resumen_objetivo=texto_resumen,
                )

            col_estado = (
                _asegurar_columna_estado(ws)
            )

            col_resumen = None

            if texto_resumen is not None:
                col_resumen = (
                    _asegurar_columna_resumen(ws)
                )

            filas_actualizadas = 0

            for fila in range(
                2,
                ws.max_row + 1,
            ):
                id_orden = _normalizar_id(
                    ws.cell(
                        row=fila,
                        column=col_id,
                    ).value
                )

                if id_orden != id_orden_buscado:
                    continue

                ws.cell(
                    row=fila,
                    column=col_estado,
                ).value = estado_rpa

                if col_resumen is not None:
                    ws.cell(
                        row=fila,
                        column=col_resumen,
                    ).value = texto_resumen

                filas_actualizadas += 1

            if filas_actualizadas == 0:
                raise ValueError(
                    f"No se encontró "
                    f"ID_ORDEN={id_orden_buscado} "
                    f"para actualizar resultado."
                )

            # ÚNICO save de la operación.
            guardar_workbook_atomico(
                wb,
                ruta_excel,
                sha256_esperado=sha256_cargado,
            )

        finally:
            wb.close()

    print(
        f"Resultado aplicado a "
        f"{filas_actualizadas} fila(s) "
        f"del ID_ORDEN={id_orden_buscado}: "
        f"ESTADO_RPA={estado_rpa}, "
        f"RESUMEN="
        f"{'actualizado' if texto_resumen is not None else 'sin cambios'}."
    )

    return filas_actualizadas



def reiniciar_orden_fallida(
    ruta_excel: str | Path,
    id_orden_buscado: Any,
) -> int:
    """
    Restablece ESTADO_RPA de 2 a 0 para un ID_ORDEN fallido.

    Reglas:
    - el ID debe existir;
    - ESTADO_RPA debe existir;
    - todas las posiciones del mismo ID deben estar en estado 2;
    - se actualizan todas las posiciones en una sola transacción;
    - RESUMEN no se modifica.
    """
    ruta_excel = Path(
        ruta_excel
    )

    id_orden_buscado = _normalizar_id(
        id_orden_buscado
    )

    if not id_orden_buscado:
        raise ValueError(
            "id_orden_buscado no puede estar vacío."
        )

    with _excel_lock(
        ruta_excel
    ):
        wb, sha256_cargado = (
            _cargar_workbook_estable_para_actualizacion(
                ruta_excel
            )
        )

        try:
            if HOJA_ORDENES not in wb.sheetnames:
                raise ValueError(
                    f"No existe la hoja '{HOJA_ORDENES}'."
                )

            ws = wb[
                HOJA_ORDENES
            ]

            headers = _headers_ws(
                ws
            )

            if COL_ID_ORDEN not in headers:
                raise ValueError(
                    f"No existe la columna "
                    f"'{COL_ID_ORDEN}'."
                )

            if COL_ESTADO_RPA not in headers:
                raise ValueError(
                    f"No existe la columna "
                    f"'{COL_ESTADO_RPA}'."
                )

            col_id = headers[
                COL_ID_ORDEN
            ]

            col_estado = headers[
                COL_ESTADO_RPA
            ]

            filas: list[int] = []
            estados: list[
                tuple[int, int | None]
            ] = []

            for fila in range(
                2,
                ws.max_row + 1,
            ):
                id_orden = _normalizar_id(
                    ws.cell(
                        row=fila,
                        column=col_id,
                    ).value
                )

                if (
                    id_orden
                    != id_orden_buscado
                ):
                    continue

                estado = _estado_a_int(
                    ws.cell(
                        row=fila,
                        column=col_estado,
                    ).value
                )

                filas.append(
                    fila
                )

                estados.append(
                    (
                        fila,
                        estado,
                    )
                )

            if not filas:
                raise ValueError(
                    f"No se encontró "
                    f"ID_ORDEN={id_orden_buscado}."
                )

            estados_invalidos = [
                (
                    fila,
                    estado,
                )
                for fila, estado
                in estados
                if estado != 2
            ]

            if estados_invalidos:
                detalle = ", ".join(
                    f"fila {fila}={estado}"
                    for fila, estado
                    in estados
                )

                raise ValueError(
                    "Solo puede reintentarse una orden "
                    "cuando todas sus posiciones tienen "
                    "ESTADO_RPA=2. "
                    f"ID_ORDEN={id_orden_buscado}. "
                    f"Estados actuales: {detalle}."
                )

            for fila in filas:
                ws.cell(
                    row=fila,
                    column=col_estado,
                ).value = 0

            guardar_workbook_atomico(
                wb,
                ruta_excel,
                sha256_esperado=(
                    sha256_cargado
                ),
            )

        finally:
            wb.close()

    print(
        f"Reintento preparado para "
        f"ID_ORDEN={id_orden_buscado}: "
        f"{len(filas)} fila(s) "
        "ESTADO_RPA 2 -> 0."
    )

    return len(
        filas
    )

def actualizar_estado_orden(
    ruta_excel: str | Path,
    id_orden_buscado: Any,
    estado_rpa: int,
) -> None:
    """
    Actualiza ESTADO_RPA en todas las filas/posiciones del ID_ORDEN.
    """
    if estado_rpa not in (0, 1, 2):
        raise ValueError("estado_rpa solo puede ser 0, 1 o 2.")

    ruta_excel = Path(ruta_excel)
    id_orden_buscado = _normalizar_id(id_orden_buscado)

    with _excel_lock(ruta_excel):
        wb, sha256_cargado = (
            _cargar_workbook_estable_para_actualizacion(
                ruta_excel
            )
        )

        try:
            if HOJA_ORDENES not in wb.sheetnames:
                raise ValueError(
                    f"No existe la hoja '{HOJA_ORDENES}'."
                )

            ws = wb[HOJA_ORDENES]
            headers = _headers_ws(ws)

            if COL_ID_ORDEN not in headers:
                raise ValueError(
                    f"No existe la columna '{COL_ID_ORDEN}'."
                )

            col_id = headers[COL_ID_ORDEN]
            col_estado = _asegurar_columna_estado(ws)

            filas_actualizadas = 0

            for fila in range(2, ws.max_row + 1):
                id_orden = _normalizar_id(
                    ws.cell(row=fila, column=col_id).value
                )

                if id_orden == id_orden_buscado:
                    ws.cell(
                        row=fila,
                        column=col_estado,
                    ).value = estado_rpa

                    filas_actualizadas += 1

            if filas_actualizadas == 0:
                raise ValueError(
                    f"No se encontró ID_ORDEN={id_orden_buscado} "
                    f"para actualizar ESTADO_RPA."
                )

            guardar_workbook_atomico(
                wb,
                ruta_excel,
                sha256_esperado=sha256_cargado,
            )

        finally:
            wb.close()

    print(
        f"ESTADO_RPA={estado_rpa} aplicado a "
        f"{filas_actualizadas} fila(s) "
        f"del ID_ORDEN={id_orden_buscado}."
    )

def crear_excel_trabajo_pendientes(
    ruta_excel: str | Path,
    *,
    solo_lectura: bool = False,
) -> tuple[Path, list[PendienteRPA]]:
    """
    Crea un Excel temporal con solo las órdenes ESTADO_RPA=0.
    Con RPA_WORKERS=2:
    - RPA_WORKER_ID=1 procesa ID_ORDEN impares.
    - RPA_WORKER_ID=2 procesa ID_ORDEN pares.
    """
    ruta_excel = Path(ruta_excel)

    if solo_lectura:
        pendientes = inspeccionar_pendientes(ruta_excel)
    else:
        pendientes = obtener_pendientes(ruta_excel)

    ids_pendientes = {p.id_orden for p in pendientes}

    directorio_temporal = obtener_directorio_temporal() or ruta_excel.parent
    directorio_temporal.mkdir(parents=True, exist_ok=True)

    ruta_trabajo = directorio_temporal / (
        f"_{ruta_excel.stem}_pendientes_RPA{_sufijo_worker()}{ruta_excel.suffix}"
    )

    hojas = pd.read_excel(ruta_excel, sheet_name=None, dtype=object)

    hojas_filtradas: dict[str, pd.DataFrame] = {}

    for nombre_hoja, df in hojas.items():
        df_filtrado = df.copy()

        # Columna técnica SOLO para el Excel temporal.
        # Sirve para que los errores indiquen la fila real del DATA.xlsx.
        if "_FILA_EXCEL_ORIGINAL" not in df_filtrado.columns:
            df_filtrado["_FILA_EXCEL_ORIGINAL"] = df_filtrado.index + 2

        if nombre_hoja == HOJA_ORDENES:
            if COL_ID_ORDEN in df_filtrado.columns:
                serie_ids = df_filtrado[COL_ID_ORDEN].map(_normalizar_id)
                df_filtrado = df_filtrado[serie_ids.isin(ids_pendientes)].copy()

        elif COL_ID_ORDEN in df_filtrado.columns:
            serie_ids = df_filtrado[COL_ID_ORDEN].map(_normalizar_id)
            df_filtrado = df_filtrado[serie_ids.isin(ids_pendientes)].copy()

        hojas_filtradas[nombre_hoja] = df_filtrado

    with pd.ExcelWriter(ruta_trabajo, engine="openpyxl") as writer:
        for nombre_hoja, df in hojas_filtradas.items():
            df.to_excel(writer, sheet_name=nombre_hoja[:31], index=False)

    return ruta_trabajo, pendientes
