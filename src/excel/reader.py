from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from src.excel.validators import (
    OrdenCotizacion,
    PosicionCotizacion,
    ValidacionExcelError,
    construir_orden_desde_grupo,
    normalizar_id,
    validar_adjunto,
)

def _fila_excel_real(fila, fila_temporal: int) -> int:
    """
    Si el Excel temporal trae _FILA_EXCEL_ORIGINAL, devuelve la fila real de DATA.xlsx.
    Si no existe, devuelve la fila temporal normal.
    """
    try:
        valor = fila.get("_FILA_EXCEL_ORIGINAL")
        if valor is not None and str(valor).strip():
            return int(float(str(valor).strip()))
    except Exception:
        pass

    return fila_temporal



HOJA_ORDENES = "Ordenes"
HOJA_ADJUNTOS = "Adjuntos"

COLUMNAS_ORDENES_REQUERIDAS = [
    "ID_ORDEN",
    "PROVEEDOR",
    "TEXTO",
    "Valor",
    "Moneda",
    "CECO",
    "Cuenta",
    "Servicio",
    "TIPO_SERVICIO",
    "TIPO_IMPUTACION",
]

COLUMNAS_ADJUNTOS = [
    "ID_ORDEN",
    "Archivo",
]


COLUMNA_TEMPLATE = "TEMPLATE"


def _es_vacio(valor: Any) -> bool:
    if valor is None:
        return True

    try:
        if pd.isna(valor):
            return True
    except Exception:
        pass

    if isinstance(valor, str) and not valor.strip():
        return True

    return False


def _normalizar_codigo(valor: Any) -> str:
    texto = str(valor).strip()

    if texto.endswith(".0") and texto[:-2].isdigit():
        texto = texto[:-2]

    return texto


def _to_decimal(valor: Any, contexto: str) -> Decimal:
    if _es_vacio(valor):
        raise ValidacionExcelError(f"{contexto}: valor vacio.")

    texto = str(valor).strip().replace(",", "")

    try:
        numero = Decimal(texto)
    except InvalidOperation as exc:
        raise ValidacionExcelError(f"{contexto}: valor no numerico '{valor}'.") from exc

    if numero <= 0:
        raise ValidacionExcelError(f"{contexto}: valor debe ser mayor a cero. Valor: {valor}")

    return numero.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)


def _fila_totalmente_vacia(fila: pd.Series, columnas: list[str]) -> bool:
    return all(_es_vacio(fila.get(columna)) for columna in columnas)


def _validar_columnas(df: pd.DataFrame, columnas_requeridas: list[str], nombre_hoja: str) -> None:
    faltantes = [col for col in columnas_requeridas if col not in df.columns]

    if faltantes:
        raise ValidacionExcelError(
            f"Faltan columnas en hoja {nombre_hoja}: " + ", ".join(faltantes)
        )


def _asegurar_columna_template(df: pd.DataFrame) -> pd.DataFrame:
    if COLUMNA_TEMPLATE not in df.columns:
        df[COLUMNA_TEMPLATE] = None

    return df



def _leer_hojas(ruta_excel: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """
    La hoja Posiciones ya no es necesaria.

    La hoja Ordenes contiene una fila por posición y puede repetir ID_ORDEN.
    La hoja Adjuntos sigue siendo opcional.
    """
    try:
        hojas = pd.read_excel(
            ruta_excel,
            sheet_name=None,
            dtype=object,
        )
    except ValueError as exc:
        raise ValidacionExcelError("No se pudo leer el archivo Excel.") from exc

    if HOJA_ORDENES not in hojas:
        raise ValidacionExcelError("El Excel debe tener la hoja 'Ordenes'.")

    df_ordenes = _asegurar_columna_template(hojas[HOJA_ORDENES].copy())

    if "TEXTO_BREVE" not in df_ordenes.columns:
        df_ordenes["TEXTO_BREVE"] = None

    if df_ordenes.empty:
        raise ValidacionExcelError("La hoja Ordenes no tiene filas de datos.")

    df_adjuntos = hojas.get(
        HOJA_ADJUNTOS,
        pd.DataFrame(columns=COLUMNAS_ADJUNTOS),
    ).copy()

    if df_adjuntos.empty:
        print("La hoja Adjuntos está vacía o no existe. Se continúa sin adjuntos.")

    _validar_columnas(df_ordenes, COLUMNAS_ORDENES_REQUERIDAS, HOJA_ORDENES)
    _validar_columnas(df_adjuntos, COLUMNAS_ADJUNTOS, HOJA_ADJUNTOS)

    return df_ordenes, df_adjuntos


def _obtener_ids_ordenes(df_ordenes: pd.DataFrame) -> set[str]:
    """
    Los ID_ORDEN repetidos están permitidos:
    cada fila representa una posición de la misma orden.
    """
    ids: set[str] = set()

    for indice, fila in df_ordenes.iterrows():
        fila_excel = _fila_excel_real(fila, indice + 2)

        try:
            id_orden = normalizar_id(fila["ID_ORDEN"])
        except ValidacionExcelError as exc:
            raise ValidacionExcelError(
                f"Error en hoja Ordenes, fila Excel {fila_excel}: {exc}"
            ) from exc

        ids.add(id_orden)

    return ids

def _agrupar_adjuntos(
    df_adjuntos: pd.DataFrame,
    ids_ordenes: set[str],
) -> dict[str, list[Path]]:
    adjuntos_por_orden: dict[str, list[Path]] = {}

    for indice, fila in df_adjuntos.iterrows():
        fila_excel = _fila_excel_real(fila, indice + 2)

        if _fila_totalmente_vacia(fila, COLUMNAS_ADJUNTOS):
            continue

        try:
            id_orden = normalizar_id(fila["ID_ORDEN"])
            ruta = validar_adjunto(fila["Archivo"])
        except ValidacionExcelError as exc:
            raise ValidacionExcelError(
                f"Error en hoja Adjuntos, fila Excel {fila_excel}: {exc}"
            ) from exc

        if id_orden not in ids_ordenes:
            raise ValidacionExcelError(
                f"Error en hoja Adjuntos, fila Excel {fila_excel}: "
                f"ID_ORDEN '{id_orden}' no existe en hoja Ordenes."
            )

        adjuntos_por_orden.setdefault(id_orden, []).append(ruta)

    for id_orden, rutas in adjuntos_por_orden.items():
        rutas_unicas = []
        vistos = set()

        for ruta in rutas:
            clave = str(ruta).lower()

            if clave not in vistos:
                vistos.add(clave)
                rutas_unicas.append(ruta)

        adjuntos_por_orden[id_orden] = rutas_unicas

    return adjuntos_por_orden


def _leer_filas_template(ruta_template: Path) -> list[tuple[str, Decimal]]:
    try:
        wb = load_workbook(ruta_template, read_only=True, data_only=True)
    except Exception as exc:
        raise ValidacionExcelError(f"No se pudo abrir TEMPLATE: {ruta_template}") from exc

    try:
        if len(wb.sheetnames) != 1:
            raise ValidacionExcelError(
                f"TEMPLATE debe tener exactamente 1 hoja. Archivo: {ruta_template}"
            )

        ws = wb[wb.sheetnames[0]]

        header_a = str(ws["A1"].value or "").strip()
        header_b = str(ws["B1"].value or "").strip()

        if header_a != "code" or header_b != "quantityOrPercent":
            raise ValidacionExcelError(
                f"TEMPLATE invalido: no modifiques cabeceras. "
                f"Se esperaba A1='code' y B1='quantityOrPercent'. Archivo: {ruta_template}"
            )

        filas: list[tuple[str, Decimal]] = []
        vistos: set[str] = set()

        for row in range(3, ws.max_row + 1):
            codigo_raw = ws.cell(row=row, column=1).value
            valor_raw = ws.cell(row=row, column=2).value

            if _es_vacio(codigo_raw) and _es_vacio(valor_raw):
                continue

            if _es_vacio(codigo_raw):
                raise ValidacionExcelError(
                    f"TEMPLATE {ruta_template}, fila {row}: codigo vacio con valor."
                )

            if _es_vacio(valor_raw):
                raise ValidacionExcelError(
                    f"TEMPLATE {ruta_template}, fila {row}: "
                    f"valor vacio para codigo {codigo_raw}."
                )

            codigo = _normalizar_codigo(codigo_raw)
            valor = _to_decimal(
                valor_raw,
                f"TEMPLATE {ruta_template}, fila {row}",
            )

            if codigo in vistos:
                raise ValidacionExcelError(
                    f"TEMPLATE {ruta_template}: codigo duplicado '{codigo}'."
                )

            vistos.add(codigo)
            filas.append((codigo, valor))

        if not filas:
            raise ValidacionExcelError(
                f"TEMPLATE no tiene filas de datos: {ruta_template}"
            )

        return filas

    finally:
        wb.close()

def _template_es_porcentaje(ruta_template: Path) -> bool:
    """
    Regla de negocio:
    Si el nombre del archivo contiene '_porcent_', se trata como template por porcentaje.
    Ejemplo: plantilla_imputacion_porcent_google.xlsx
    """
    nombre = str(ruta_template.name).strip().lower()
    return "_porcent_" in nombre

def _validar_template_imputacion(
    ruta_template: Path,
    ceco_cabecera: str,
    valor_posicion: Decimal,
    contexto: str,
) -> None:
    filas = _leer_filas_template(ruta_template)

    primer_codigo = filas[0][0]

    if primer_codigo != ceco_cabecera:
        raise ValidacionExcelError(
            f"{contexto}: el primer codigo del TEMPLATE debe ser el CECO/imputacion "
            f"de cabecera '{ceco_cabecera}', pero se encontro '{primer_codigo}'. "
            f"Archivo: {ruta_template}"
        )

    suma = sum((valor for _, valor in filas), Decimal("0")).quantize(Decimal("0.0001"))

    if _template_es_porcentaje(ruta_template):
        sumas_validas_porcentaje = {
            Decimal("1.0000"),
            Decimal("100.0000"),
        }

        if suma not in sumas_validas_porcentaje:
            raise ValidacionExcelError(
                f"{contexto}: TEMPLATE marcado como PORCENTAJE por nombre de archivo '_porcent_', "
                f"pero la suma es {suma}. Debe sumar 1.0000 o 100.0000. "
                f"Archivo: {ruta_template}"
            )

        print(
            f"Info: {contexto}: TEMPLATE detectado como PORCENTAJE por nombre de archivo. "
            f"Suma validada: {suma}. Archivo: {ruta_template}"
        )
        return

    if suma != valor_posicion:
        raise ValidacionExcelError(
            f"{contexto}: la suma del TEMPLATE es {suma}, "
            f"pero el valor de la posicion es {valor_posicion}. "
            f"Archivo: {ruta_template}"
        )



def _validar_posiciones_de_orden(orden: OrdenCotizacion) -> None:
    """
    En el modelo nuevo todas las posiciones vienen de las filas agrupadas
    de la hoja Ordenes.
    """
    if not orden.posiciones_detalle:
        raise ValidacionExcelError(
            f"La orden '{orden.id_orden}' no tiene posiciones válidas en la hoja Ordenes."
        )

    if len(orden.posiciones_detalle) != orden.posiciones:
        raise ValidacionExcelError(
            f"La orden '{orden.id_orden}' indica internamente {orden.posiciones} posiciones, "
            f"pero se construyeron {len(orden.posiciones_detalle)}."
        )

    suma_posiciones = sum(
        (posicion.valor for posicion in orden.posiciones_detalle),
        Decimal("0"),
    ).quantize(Decimal("0.0001"))

    if suma_posiciones != orden.valor:
        raise ValidacionExcelError(
            f"La orden '{orden.id_orden}' tiene total calculado={orden.valor}, "
            f"pero la suma de sus posiciones es {suma_posiciones}."
        )

    for indice, posicion in enumerate(orden.posiciones_detalle, start=1):
        if posicion.template:
            _validar_template_imputacion(
                ruta_template=posicion.template,
                ceco_cabecera=posicion.ceco,
                valor_posicion=posicion.valor,
                contexto=f"Orden {orden.id_orden}, posicion {indice}",
            )


def leer_ordenes_excel(ruta_excel: str | Path) -> list[OrdenCotizacion]:
    ruta_excel = Path(ruta_excel)

    if not ruta_excel.exists():
        raise FileNotFoundError(f"No existe el archivo Excel: {ruta_excel}")

    df_ordenes, df_adjuntos = _leer_hojas(ruta_excel)

    ids_ordenes = _obtener_ids_ordenes(df_ordenes)
    adjuntos_por_orden = _agrupar_adjuntos(df_adjuntos, ids_ordenes)

    ordenes: list[OrdenCotizacion] = []

    # sort=False conserva el orden en que aparecen los ID_ORDEN en el Excel.
    for id_orden_raw, grupo in df_ordenes.groupby("ID_ORDEN", sort=False, dropna=False):
        primera_fila = grupo.iloc[0]
        fila_excel = _fila_excel_real(primera_fila, int(grupo.index[0]) + 2)

        try:
            id_orden = normalizar_id(id_orden_raw)
            adjuntos = adjuntos_por_orden.get(id_orden, [])

            orden = construir_orden_desde_grupo(
                grupo=grupo,
                adjuntos=adjuntos,
            )

            _validar_posiciones_de_orden(orden)
            ordenes.append(orden)

            filas_reales = [
                _fila_excel_real(fila, int(indice) + 2)
                for indice, fila in grupo.iterrows()
            ]

            print(
                f"Orden {id_orden}: {len(grupo)} posicion(es) agrupadas "
                f"desde filas Excel {filas_reales}. Total={orden.valor}"
            )

        except ValidacionExcelError as exc:
            raise ValidacionExcelError(
                f"Error en hoja Ordenes, grupo ID_ORDEN={id_orden_raw}, "
                f"primera fila Excel {fila_excel}: {exc}"
            ) from exc

    return ordenes

def leer_primera_orden_excel(ruta_excel: str | Path) -> OrdenCotizacion:
    ordenes = leer_ordenes_excel(ruta_excel)

    if not ordenes:
        raise ValidacionExcelError("No se encontro ninguna orden valida en el Excel.")

    return ordenes[0]
