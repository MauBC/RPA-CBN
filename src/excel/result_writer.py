from __future__ import annotations

from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from typing import Any
import os
import time

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from src.utils.app_paths import obtener_directorio_salida_compatibilidad




def _asegurar_output_dir() -> Path:
    ruta = obtener_directorio_salida_compatibilidad()
    ruta.mkdir(parents=True, exist_ok=True)
    return ruta


def _ruta_resultado(ruta_excel: str | Path) -> Path:
    ruta_excel = Path(ruta_excel)
    worker_id = os.getenv("RPA_WORKER_ID", "").strip()
    sufijo = f"_worker_{worker_id}" if worker_id else ""

    output_dir = _asegurar_output_dir()
    return output_dir / f"resultado_{ruta_excel.stem}{sufijo}.xlsx"


@contextmanager
def _excel_lock(ruta_excel: str | Path, timeout_segundos: int = 180):
    ruta_excel = Path(ruta_excel)
    ruta_lock = Path(str(ruta_excel) + ".lock")
    inicio = time.time()
    fd = None

    while True:
        try:
            fd = os.open(str(ruta_lock), os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.write(fd, f"pid={os.getpid()} time={time.time()}".encode("utf-8"))
            break

        except FileExistsError:
            try:
                if ruta_lock.exists() and time.time() - ruta_lock.stat().st_mtime > 600:
                    ruta_lock.unlink()
                    continue
            except Exception:
                pass

            if time.time() - inicio > timeout_segundos:
                raise TimeoutError(f"No se pudo obtener lock de Excel: {ruta_lock}")

            time.sleep(0.5)

    try:
        yield

    finally:
        if fd is not None:
            try:
                os.close(fd)
            except Exception:
                pass

        try:
            ruta_lock.unlink()
        except FileNotFoundError:
            pass
        except Exception:
            pass


def _ajustar_hoja_resultado(
    ruta_salida: Path,
    hoja: str = "Resultado_Ordenes",
) -> None:
    wb = load_workbook(ruta_salida)

    try:
        if hoja not in wb.sheetnames:
            wb.save(ruta_salida)
            return

        ws = wb[hoja]

        fill_header = PatternFill(
            "solid",
            fgColor="1F4E78",
        )

        font_header = Font(
            color="FFFFFF",
            bold=True,
        )

        for cell in ws[1]:
            cell.fill = fill_header
            cell.font = font_header
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for col_idx, column_cells in enumerate(
            ws.columns,
            start=1,
        ):
            max_len = 10

            for cell in column_cells:
                value = (
                    ""
                    if cell.value is None
                    else str(cell.value)
                )

                max_len = max(
                    max_len,
                    min(len(value) + 2, 70),
                )

            ws.column_dimensions[
                get_column_letter(col_idx)
            ].width = max_len

        for row in ws.iter_rows(min_row=2):
            estado = (
                str(row[3].value or "").upper()
                if len(row) >= 4
                else ""
            )

            if estado == "OK":
                fill = PatternFill(
                    "solid",
                    fgColor="E2F0D9",
                )

            elif estado.startswith("ERROR"):
                fill = PatternFill(
                    "solid",
                    fgColor="FCE4D6",
                )

            else:
                fill = PatternFill(
                    "solid",
                    fgColor="FFF2CC",
                )

            for cell in row:
                cell.fill = fill
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        ws.freeze_panes = "A2"

        wb.save(ruta_salida)

    finally:
        wb.close()

def guardar_resultados_ordenes(
    ruta_excel: str | Path,
    resultados: list[dict[str, Any]],
) -> Path:
    _asegurar_output_dir()

    ruta_excel = Path(ruta_excel)
    ruta_salida = _ruta_resultado(ruta_excel)

    df_resultado = pd.DataFrame(resultados)

    columnas = [
        "NRO",
        "ID_ORDEN",
        "FILA_EXCEL",
        "ESTADO",
        "N_COTIZACION",
        "ESTADO_CBN",
        "MENSAJE_USUARIO",
        "PASO_ERROR",
        "DETALLE_TECNICO",
        "LOG",
        "SCREENSHOT",
        "HTML",
        "FECHA_HORA",
    ]

    for columna in columnas:
        if columna not in df_resultado.columns:
            df_resultado[columna] = ""

    df_resultado = df_resultado[columnas]

    try:
        with _excel_lock(ruta_excel):
            hojas = pd.read_excel(ruta_excel, sheet_name=None, dtype=object)
    except Exception:
        hojas = {}

    with pd.ExcelWriter(ruta_salida, engine="openpyxl") as writer:
        for nombre_hoja, df in hojas.items():
            nombre_limpio = str(nombre_hoja)[:31]
            df.to_excel(writer, sheet_name=nombre_limpio, index=False)

        df_resultado.to_excel(writer, sheet_name="Resultado_Ordenes", index=False)

    _ajustar_hoja_resultado(ruta_salida)

    return ruta_salida


def guardar_resultado_error_validacion(
    ruta_excel: str | Path,
    mensaje_error: str,
) -> Path:
    resultado = [
        {
            "NRO": 1,
            "ID_ORDEN": "",
            "ESTADO": "ERROR_VALIDACION_EXCEL",
            "MENSAJE_USUARIO": mensaje_error,
            "PASO_ERROR": "validacion_excel",
            "DETALLE_TECNICO": mensaje_error,
            "LOG": "",
            "SCREENSHOT": "",
            "HTML": "",
            "FECHA_HORA": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }
    ]

    return guardar_resultados_ordenes(ruta_excel, resultado)
