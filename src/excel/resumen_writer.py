from __future__ import annotations

from contextlib import contextmanager
from pathlib import Path
from typing import Any
import os
import time

import pandas as pd
from openpyxl import load_workbook


HOJA_ORDENES = "Ordenes"
COL_ID_ORDEN = "ID_ORDEN"
COL_RESUMEN = "RESUMEN"


def _normalizar_header(valor: Any) -> str:
    return str(valor or "").strip().upper()


def _normalizar_id(valor: Any) -> str:
    if valor is None:
        return ""

    try:
        if pd.isna(valor):
            return ""
    except Exception:
        pass

    texto = str(valor).strip()

    if texto.endswith(".0"):
        texto = texto[:-2]

    return texto


@contextmanager
def _excel_lock(ruta_excel: str | Path, timeout_segundos: int = 180):
    ruta_excel = Path(ruta_excel)
    ruta_lock = Path(str(ruta_excel) + ".lock")
    inicio = time.time()
    fd = None

    while True:
        try:
            fd = os.open(str(ruta_lock), os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.write(fd, f"pid={os.getpid()} time={time.time()}".encode("utf-8"))
            break

        except FileExistsError:
            try:
                if ruta_lock.exists() and time.time() - ruta_lock.stat().st_mtime > 600:
                    ruta_lock.unlink()
                    continue
            except Exception:
                pass

            if time.time() - inicio > timeout_segundos:
                raise TimeoutError(f"No se pudo obtener lock de Excel: {ruta_lock}")

            time.sleep(0.5)

    try:
        yield

    finally:
        if fd is not None:
            try:
                os.close(fd)
            except Exception:
                pass

        try:
            ruta_lock.unlink()
        except FileNotFoundError:
            pass
        except Exception:
            pass


def _headers_ws(ws) -> dict[str, int]:
    headers = {}

    for cell in ws[1]:
        nombre = _normalizar_header(cell.value)

        if nombre:
            headers[nombre] = cell.column

    return headers


def _asegurar_columna_resumen(ws) -> int:
    headers = _headers_ws(ws)

    if COL_RESUMEN in headers:
        return headers[COL_RESUMEN]

    nueva_col = ws.max_column + 1
    ws.cell(row=1, column=nueva_col).value = COL_RESUMEN

    return nueva_col



def actualizar_resumen_orden(
    ruta_excel: str | Path,
    id_orden_buscado: Any,
    resumen: str,
) -> None:
    """
    Actualiza RESUMEN en todas las filas/posiciones del ID_ORDEN.
    """
    ruta_excel = Path(ruta_excel)
    id_orden_buscado = _normalizar_id(id_orden_buscado)
    texto_resumen = str(resumen or "").strip()

    with _excel_lock(ruta_excel):
        wb = load_workbook(ruta_excel)

        if HOJA_ORDENES not in wb.sheetnames:
            raise ValueError(f"No existe la hoja '{HOJA_ORDENES}'.")

        ws = wb[HOJA_ORDENES]
        headers = _headers_ws(ws)

        if COL_ID_ORDEN not in headers:
            raise ValueError(f"No existe la columna '{COL_ID_ORDEN}'.")

        col_id = headers[COL_ID_ORDEN]
        col_resumen = _asegurar_columna_resumen(ws)

        filas_actualizadas = 0

        for fila in range(2, ws.max_row + 1):
            id_orden = _normalizar_id(ws.cell(row=fila, column=col_id).value)

            if id_orden == id_orden_buscado:
                ws.cell(row=fila, column=col_resumen).value = texto_resumen
                filas_actualizadas += 1

        if filas_actualizadas == 0:
            raise ValueError(
                f"No se encontró ID_ORDEN={id_orden_buscado} para actualizar RESUMEN."
            )

        wb.save(ruta_excel)

    print(
        f"RESUMEN aplicado a {filas_actualizadas} fila(s) "
        f"del ID_ORDEN={id_orden_buscado}."
    )

